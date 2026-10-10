from pathlib import Path
from docx import Document

src = Path(r"C:\Users\admin\Downloads\临床决策 App 竞品分析 (1).docx")
out = Path(r"C:\Users\admin\Documents\ChatGPT\大查房\竞品分析_临床决策\临床决策_App_竞品分析_精简版.docx")

for label, path in [("source", src), ("output", out)]:
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    print(label, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "chars", len(text), "images", len(doc.inline_shapes))

doc = Document(out)
needles = (
    "Step 4：", "Step 4-A：", "Step 4-C：", "Step 4-D：", "Step 5：", "Step 5-A：", "Step 5-B：",
    "Step 6：", "Step 7：", "Step 8-A：", "Step 8-B：", "Step 9：", "Step 10：",
    "五、", "P0-", "P1-", "P2-", "七、", "商业策略", "具体行动", "八、", "证据可见性", "九、",
)
for i, p in enumerate(doc.paragraphs):
    if any(p.text.startswith(n) for n in needles) or p.text.startswith("证据：") or "详见 P" in p.text:
        print(f"P{i:03d} {p.text}")

for ti, table in enumerate(doc.tables):
    first = " | ".join(c.text.replace("\n", " / ") for c in table.rows[0].cells)
    print(f"T{ti} rows={len(table.rows)} cols={len(table.columns)} first={first}")
    if first == "优先级 | 建议 | 目标":
        for row in table.rows:
            print("  ", " | ".join(c.text for c in row.cells))
