from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(r"C:\Users\admin\Documents\ChatGPT\大查房\deliverables\OncoRounds-最新源码代码修改说明-2026-08-24.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "17243A"
MUTED = "5F6C84"
LIGHT = "E8EEF5"
VERY_LIGHT = "F4F6F9"
BORDER = "D7E0EA"
WHITE = "FFFFFF"
CODE_BG = "F5F7FA"
GREEN = "18796A"


def set_run_font(run, ascii_font="Calibri", east_asia="Microsoft YaHei", size=None,
                 color=None, bold=None, italic=None):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_numbering(doc, kind="bullet"):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    # OOXML requires every abstractNum before the concrete num entries.
    # Inserting at the end makes Word repair/merge the lists, which can turn
    # bullets into continued decimal numbering after rendering.
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header_p = section.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
header_run = header_p.add_run("OncoRounds  ·  技术交接说明")
set_run_font(header_run, size=9, color=MUTED, bold=True)

footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_run = footer_p.add_run("2026-08-24  |  第 ")
set_run_font(footer_run, size=9, color=MUTED)
add_page_number(footer_p)
footer_run2 = footer_p.add_run(" 页")
set_run_font(footer_run2, size=9, color=MUTED)

bullet_num = add_numbering(doc, "bullet")
decimal_num = add_numbering(doc, "decimal")


def add_body(text, bold_prefix=None, color=INK, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, color=color, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=color)
    return p


def add_bullet(text, num_id=None):
    p = doc.add_paragraph()
    apply_num(p, bullet_num if num_id is None else num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_step(text, num_id=None):
    p = doc.add_paragraph()
    apply_num(p, decimal_num if num_id is None else num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_code(code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CODE_BG)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "14")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)
    for index, line in enumerate(code.splitlines()):
        if index:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, ascii_font="Consolas", east_asia="Microsoft YaHei", size=9.2, color="26384F")
    return p


def add_callout(label, text, fill=LIGHT):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=11, color=DARK_BLUE, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color=INK)
    return p


# First-page masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("ONCOROUNDS · SOURCE HANDOFF")
set_run_font(r, size=10, color=BLUE, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
r = p.add_run("OncoRounds 最新源码代码修改说明")
set_run_font(r, size=25, color=INK, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("最终 Demo 同步、答题进度字段与 TimePoint 数据映射")
set_run_font(r, size=13, color=MUTED)

for label, value in (
    ("版本", "2026-08-24"),
    ("适用源码", "OncoRounds-最新完整源码-2026-08-24"),
    ("目标", "确保完整源码重新构建后与最终离线 Demo 行为一致"),
    ("修改边界", "不修改病例内容、三套5题题库、答案或 FastGPT 接口"),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(label + "：")
    set_run_font(r1, size=10.5, color=DARK_BLUE, bold=True)
    r2 = p.add_run(value)
    set_run_font(r2, size=10.5, color=INK)

add_callout(
    "关键结论",
    "本次使用最小无损同步方案。src 目录没有业务代码改动；最终 Demo 的增强 CSS、JS 和疗效图片被纳入 public/assets，并由两个入口文件稳定加载。",
)

doc.add_heading("1. 修改范围总览", level=1)
add_body("以下内容是相对于 2026-08-19 完整源码的实际差异。dist 属于重新构建产物，不应手工维护。")

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
set_table_geometry(table, [720, 3240, 1260, 4140])
headers = ["序号", "文件", "类型", "作用"]
for i, value in enumerate(headers):
    cell = table.rows[0].cells[i]
    set_cell_shading(cell, LIGHT)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(value)
    set_run_font(r, size=9.5, color=DARK_BLUE, bold=True)
set_repeat_table_header(table.rows[0])

rows = [
    ("1", "index.html", "修改", "开发入口加载增强 CSS 与 JS。"),
    ("2", "vite.index.html", "修改", "生产构建前覆盖 index.html，必须同步修改。"),
    ("3", "public/assets/lung-treatment-response.css", "新增", "移动端布局、真实 LOGO、疗效卡片和病例回顾弹层样式。"),
    ("4", "public/assets/lung-treatment-response.js", "新增", "注入疗效卡片、缓存病例摘要/病程轴、提供答题页回顾弹层。"),
    ("5", "public/assets/cases/lung-treatment-response.png", "新增", "ROS1 疗效影像评估图片。"),
    ("6", "dist/*", "构建生成", "pnpm build 生成的可部署产物，不直接编辑。"),
]
for row_values in rows:
    cells = table.add_row().cells
    for idx, value in enumerate(row_values):
        p = cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(value)
        set_run_font(r, size=9.2, color=INK)

doc.add_heading("2. 入口文件如何修改", level=1)
doc.add_heading("2.1 index.html", level=2)
add_body("修改位置：head 内 description 后，以及 body 结束前。")
add_code('''<meta name="description" content="肿瘤医生大查房互动教学 Demo" />
<link rel="stylesheet" href="./assets/lung-treatment-response.css?v=20260824-1" />
<title>肿瘤医生大查房</title>

...

<script type="module" src="/src/main.tsx"></script>
<script src="./assets/lung-treatment-response.js?v=20260824-1"></script>''')

doc.add_heading("2.2 vite.index.html", level=2)
add_body("使用完全相同的两行引用。package.json 的 prebuild 会执行以下命令：")
add_code('''node -e "require('node:fs').copyFileSync('vite.index.html', 'index.html')"''')
add_callout("为什么两个文件都要改", "如果只改 index.html，下一次 pnpm build 会被 vite.index.html 覆盖，增强样式和交互会再次消失。", VERY_LIGHT)

doc.add_heading("3. 增强 JavaScript 的具体职责", level=1)
add_body("文件：public/assets/lung-treatment-response.js。它保持原 React 业务逻辑不变，通过监听页面切换补充最终 Demo 已存在的功能。")

for title, detail in (
    ("syncResponseCard()（约第16行）", "仅在 ROS1 病例摘要中插入“疗效影像评估”卡片，并引用 lung-treatment-response.png。"),
    ("cacheVisibleOverview()（约第63行）", "缓存当前病例已渲染的病例摘要与病程轴 DOM，供答题时回顾。"),
    ("hydrateCaseContent()（约第86行）", "必要时自动切换摘要/病程轴页签取得两份内容，然后恢复用户原页签。"),
    ("openReviewLayer()（约第154行）", "在答题页创建病例资料对话框，包含“病例摘要”和“病程轴”两个页签。"),
    ("ensureReviewTrigger()（约第204行）", "把答题页右侧状态点替换成“回顾病例资料”按钮。"),
    ("MutationObserver（约第239行）", "监听 React 根节点变化，在页面切换后重新同步增强内容。"),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r1 = p.add_run(title + "：")
    set_run_font(r1, size=11, color=DARK_BLUE, bold=True)
    r2 = p.add_run(detail)
    set_run_font(r2, size=11, color=INK)

add_code('''new MutationObserver(syncAppEnhancements).observe(
  document.getElementById("root"),
  { childList: true, subtree: true }
);''')

doc.add_heading("4. 增强 CSS 的具体职责", level=1)
add_body("文件：public/assets/lung-treatment-response.css。主要选择器和媒体查询如下。")
for text in (
    "@media (min-height: 700px)：调整高屏设备首页纵向重心。",
    "@media (max-width: 559px)：修复手机端顶部空间、主图高度和底部操作区。",
    ".brand-lockup.is-compact .brand-orb：用真实 oncorounds-orb.png 替代原图标。",
    ".treatment-response-card：定义 ROS1 疗效影像卡片。",
    ".case-review-trigger：定义“回顾病例资料”紧凑胶囊按钮。",
    ".case-review-layer：定义全屏病例回顾弹层。",
    ".case-review-tabs / .case-review-scroll：定义摘要/病程轴切换和滚动区域。",
):
    add_bullet(text)

doc.add_heading("5. 答题页顶部进度的数据来源", level=1)
add_body("页面位置：src/components/RoundSession.tsx。顶部病例简称来自 clinicalCase.shortTitle；题号和进度由 questionIndex 与 questions.length 实时计算。")
add_code('''const progress =
  ((state.questionIndex + 1) / clinicalCase.questions.length) * 100

<span>{clinicalCase.shortTitle}</span>
<strong>
  {state.questionIndex + 1} / {clinicalCase.questions.length}
</strong>

<div className="progress-track"
     aria-label={`查房进度 ${Math.round(progress)}%`}>
  <span style={{ width: `${progress}%` }} />
</div>''')

add_body("独立交付文件：01-答题进度.json。每个病例有5个 steps，进度值固定为20、40、60、80、100。")
add_code('''{
  "case_id": "lung-ros1",
  "display_title": "ROS1 融合肺癌",
  "total_questions": 5,
  "steps": [
    {
      "question_id": "Q2",
      "current_question": 2,
      "total_questions": 5,
      "display_counter": "2 / 5",
      "progress_percent": 40
    }
  ]
}''')

doc.add_heading("6. TimePoint 胶囊的数据来源", level=1)
add_body("原始字段来自已审核题库中的 question.timepoint 和 question.phase。src/data/publishedQuestions.ts 将两者按“空格·空格”连接为 question.context。")
add_code('''context: [question.timepoint, question.phase]
  .filter(Boolean)
  .join(' · ')''')
add_body("RoundSession.tsx 只负责显示已经转换好的 context：")
add_code('''<div className="question-context">
  <ClipboardText weight="duotone" />
  <p>{question.context}</p>
</div>''')
add_body("独立交付文件：02-题目时间点胶囊.json。每道题保留 timepoint、phase、display_text 和对应题干。")
add_code('''{
  "question_id": "Q2",
  "timepoint": "2021 年 9 月 CNS 进展时",
  "phase": "疾病进展管理",
  "display_text": "2021 年 9 月 CNS 进展时 · 疾病进展管理"
}''')

doc.add_heading("7. 构建与交付步骤", level=1)
build_steps_num = add_numbering(doc, "decimal")
for text in (
    "在 OncoRounds-最新完整源码-2026-08-24 目录打开终端。",
    "执行 pnpm install --frozen-lockfile 安装锁定依赖。",
    "执行 pnpm test -- --run 运行测试。",
    "执行 pnpm build 进行 TypeScript 检查和 Vite 生产构建。",
    "执行 pnpm preview 打开 dist 进行最终浏览器检查。",
):
    add_step(text, build_steps_num)
add_code('''pnpm install --frozen-lockfile
pnpm test -- --run
pnpm build
pnpm preview''')

doc.add_heading("8. 验证结果", level=1)
verification_bullets_num = add_numbering(doc, "bullet")
for text in (
    "3个测试文件、11项测试全部通过。",
    "TypeScript 独立检查通过。",
    "Vite 生产构建通过。",
    "dist 中包含增强 CSS、增强 JS 和疗效图片。",
    "实际浏览器验证：首页轮播、ROS1疗效影像、答题页病例回顾均正常。",
    "LOGO与病例图片加载成功，控制台无错误。",
    "src 目录文件哈希对比：业务代码改动数量为0。",
):
    add_bullet(text, verification_bullets_num)

doc.add_heading("9. 维护注意事项", level=1)
maintenance_bullets_num = add_numbering(doc, "bullet")
for text in (
    "不要直接修改 dist；dist 必须由 pnpm build 重新生成。",
    "不要删除 index.html 与 vite.index.html 中的增强 CSS/JS 引用。",
    "不要只更新 index.html；prebuild 会用 vite.index.html 覆盖它。",
    "不要把 FastGPT API Key 写入前端源码、JSON、日志或 Git 仓库。",
    "题库仍执行严格5题和发布审核门禁；本次修改没有改变题目内容与答案。",
    "如后续要移除 DOM 增强层，应先把病例回顾和疗效影像正式迁移为 React 组件，再删除 public/assets 中对应 JS/CSS。",
):
    add_bullet(text, maintenance_bullets_num)

add_callout("交付状态", "最新版完整源码、进度 JSON、TimePoint JSON 均已单独交付；此文档用于开发复现、审查和后续维护。", "E9F5F2")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
