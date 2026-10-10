from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\admin\Documents\ChatGPT\大查房")
OUT = ROOT / "outputs" / "中医问答工作流详细总结.docx"
QA_DIR = ROOT / "work" / "docx_qa"
FLOW_IMG = QA_DIR / "tcm_workflow_overview.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "7A5A00"
RED = "9B1C1C"
GREEN = "2E6F40"
WHITE = "FFFFFF"
BLACK = "222222"
MUTED = "666666"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)

    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)

    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, bold=None, color=None, italic=None,
                 ascii_font="Calibri", east_asia="Microsoft YaHei"):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph_runs(paragraph, size=11, color=BLACK, bold=False):
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bold)


def set_repeatable_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.188)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25

    if "Small Text" not in styles:
        st = styles.add_style("Small Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        st = styles["Small Text"]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(9)
    st.font.color.rgb = RGBColor.from_string(MUTED)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15


def add_field(run, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, text, end])


def setup_header_footer(section):
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("中医问答工作流学习报告")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    r2 = p.add_run("  |  基于实际节点配置")
    set_run_font(r2, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run("第 ")
    set_run_font(r, size=8.5, color=MUTED)
    r = fp.add_run()
    set_run_font(r, size=8.5, color=MUTED)
    add_field(r, "PAGE")
    r = fp.add_run(" 页 / 共 ")
    set_run_font(r, size=8.5, color=MUTED)
    r = fp.add_run()
    set_run_font(r, size=8.5, color=MUTED)
    add_field(r, "NUMPAGES")
    r = fp.add_run(" 页")
    set_run_font(r, size=8.5, color=MUTED)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    r = p.add_run(text)
    set_run_font(r, size=11, color=BLACK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_run_font(r, size=11, color=BLACK)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=11, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=11, color=BLACK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=BLACK)
    return p


def add_callout(doc, label, text, fill=CALLOUT, accent=BLUE, trailing_spacer=True):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + "  ")
    set_run_font(r, size=10.5, bold=True, color=accent)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=BLACK)
    if trailing_spacer:
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, widths_dxa, font_size=9.5, zebra=True):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    table.style = "Table Grid"
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True, color=INK)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        if zebra and ridx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, "FAFBFC")
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def wrap_cn(text, max_chars):
    lines = []
    line = ""
    for ch in text:
        line += ch
        if len(line) >= max_chars or ch == "\n":
            lines.append(line.strip())
            line = ""
    if line:
        lines.append(line)
    return "\n".join(lines)


def create_flowchart(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1600, 1880
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    if not font_path.exists():
        font_path = Path(r"C:\Windows\Fonts\simhei.ttf")
    title_font = ImageFont.truetype(str(font_path), 44)
    box_font = ImageFont.truetype(str(font_path), 30)
    small_font = ImageFont.truetype(str(font_path), 24)

    d.text((width // 2, 35), "中医问答工作流：主链路与关键分支", font=title_font,
           fill=(31, 77, 120), anchor="ma")

    def box(x, y, w, h, text, fill="#E8EEF5", outline="#2E74B5", font=box_font):
        d.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=fill,
                            outline=outline, width=4)
        wrapped = wrap_cn(text, max(8, int(w / 34)))
        d.multiline_text((x + w / 2, y + h / 2), wrapped, font=font,
                         fill=(24, 38, 56), anchor="mm", align="center", spacing=8)
        return (x, y, w, h)

    def arrow(p1, p2, color="#5B83B4", width_px=5):
        d.line([p1, p2], fill=color, width=width_px)
        import math
        ang = math.atan2(p2[1] - p1[1], p2[0] - p1[0])
        size = 17
        pts = [
            p2,
            (p2[0] - size * math.cos(ang - 0.55), p2[1] - size * math.sin(ang - 0.55)),
            (p2[0] - size * math.cos(ang + 0.55), p2[1] - size * math.sin(ang + 0.55)),
        ]
        d.polygon(pts, fill=color)

    b1 = box(500, 120, 600, 105, "流程开始：用户问题、文件、全局变量")
    b2 = box(500, 270, 600, 115, "附件标准化 + 简易/完整上下文")
    b3 = box(500, 430, 600, 105, "意图分析 + 过滤")
    arrow((800, 225), (800, 270))
    arrow((800, 385), (800, 430))

    left = box(95, 610, 560, 125, "IF：有文件\n通过模型知识回答", fill="#FFF4E5", outline="#B7791F")
    right = box(945, 610, 560, 125, "无文件：意图转数组\n进入并行执行", fill="#EAF6EE", outline="#2E6F40")
    arrow((800, 535), (375, 610))
    arrow((800, 535), (1225, 610))

    b4 = box(945, 790, 560, 135, "每个意图：清理上下文\n构造检索词 + 选择文献类型", fill="#EAF6EE", outline="#2E6F40")
    b5 = box(945, 980, 560, 115, "TCM文献检索与重排", fill="#EAF6EE", outline="#2E6F40")
    b6 = box(945, 1150, 560, 120, "成功结果拍平并跨路去重\n按相关分数排序、重新编号", fill="#EAF6EE", outline="#2E6F40", font=small_font)
    arrow((1225, 735), (1225, 790))
    arrow((1225, 925), (1225, 980))
    arrow((1225, 1095), (1225, 1150))

    export_box = box(70, 1125, 520, 135, "证据导出支线：删除内部字段\n保存检索证据并生成访问URL", fill="#F4F6F9", outline="#6B7280", font=small_font)
    arrow((945, 1210), (590, 1192), color="#6B7280")

    b7 = box(520, 1340, 560, 120, "引用模板 + 通过知识库回答", fill="#E8EEF5", outline="#2E74B5")
    arrow((1225, 1270), (800, 1340))
    b8 = box(520, 1510, 560, 115, "判断是否是“有效回答”", fill="#FDECEC", outline="#9B1C1C")
    arrow((800, 1460), (800, 1510))

    normal = box(110, 1690, 500, 105, "正常：写入最终答案变量", fill="#EAF6EE", outline="#2E6F40", font=small_font)
    fallback = box(990, 1680, 500, 125, "证据不足：标记二次回答\n回退到模型自身知识", fill="#FDECEC", outline="#9B1C1C", font=small_font)
    arrow((800, 1625), (360, 1690), color="#2E6F40")
    arrow((800, 1625), (1240, 1680), color="#9B1C1C")

    # final convergence label under the diagram
    d.text((800, 1840), "两条回答路径汇合 → 格式整理 → 输出给用户",
           font=small_font, fill=(11, 37, 69), anchor="mm")
    img.save(path, quality=95)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(78)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WORKFLOW STUDY GUIDE")
    set_run_font(r, size=10, bold=True, color=GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("中医问答工作流")
    set_run_font(r, size=30, bold=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("详细拆解与风险评估")
    set_run_font(r, size=22, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(52)
    r = p.add_run("基于实际截图、节点配置、代码与提示词")
    set_run_font(r, size=12.5, color=MUTED)

    add_callout(
        doc,
        "阅读边界",
        "本文只记录已提供配置能够证明的事实；看不清、未提供或无法由运行轨迹确认的内容均标记为“待确认”。",
        fill="F8F6F0",
        accent=GOLD,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(52)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("版本日期：2026年8月18日")
    set_run_font(r, size=10.5, color=MUTED, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("用途：工作流学习与设计复盘，不代表临床有效性评价")
    set_run_font(r, size=9.5, color=MUTED, italic=True)
    doc.add_page_break()


def build_document():
    QA_DIR.mkdir(parents=True, exist_ok=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    create_flowchart(FLOW_IMG)

    doc = Document()
    doc.core_properties.title = "中医问答工作流详细拆解与风险评估"
    doc.core_properties.subject = "基于实际节点配置的工作流学习报告"
    doc.core_properties.author = "OpenAI Codex"
    doc.core_properties.keywords = "中医问答, FastGPT, RAG, 工作流, 检索, 风险评估"

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    setup_header_footer(section)
    set_repeatable_styles(doc)
    add_cover(doc)

    doc.add_heading("执行摘要", level=1)
    add_callout(
        doc,
        "一句话结论",
        "这是一套以“意图分类—多路中医文献检索—证据排序—循证回答—证据不足时模型知识回退—最终格式整理”为主线的中医问答工作流，并将检索证据另存到对象存储。",
        fill=LIGHT_BLUE,
        accent=INK,
    )
    add_body(doc, "其前半段已经具备较完整的RAG工程结构：附件标准化、意图拆分、并行检索、文献类型过滤、Rerank、跨路去重和引用模板。其后半段存在一个决定性的安全矛盾：知识库证据不足时并未真正停止，而是回退到模型自身知识继续生成医疗回答；最终格式节点还会删除正文引用角标。")
    add_body(doc, "因此，这套流程适合学习多意图RAG的结构与工程分层，但从当前配置看，不能把最终输出视为经过严格证据校验、适合直接临床使用的中医决策结果。")

    doc.add_heading("完整结构图", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    flow_shape = run.add_picture(str(FLOW_IMG), width=Inches(6.25))
    flow_shape._inline.docPr.set(
        "descr",
        "中医问答工作流主链路：输入与附件标准化后进行意图分析；有文件时走模型知识回答，无文件时按意图并行检索中医文献、重排证据并生成知识库回答；证据不足分支会回退模型知识；两条路径最后统一格式化输出。",
    )
    flow_shape._inline.docPr.set("title", "中医问答工作流主链路与关键分支")
    p.paragraph_format.space_after = Pt(4)
    cap = doc.add_paragraph(style="Small Text")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run("图1  中医问答工作流主链路。红色分支表示证据不足后仍回退到模型自身知识的高风险路径。")

    doc.add_heading("一、节点清单与作用", level=1)
    node_rows = [
        ("1", "流程开始", "接收用户问题和文件链接，并提供患者摘要、附件上下文、回答状态、会话ID和当前时间等全局变量。"),
        ("2", "拼接用户用户附件上下文", "把附件对象整理为完整result与最多300字符的truncatedResult；OCR为空时插入能力边界提示。"),
        ("3", "前置-用户问题拼接成简易上下文", "合并用户问题、patientSummarySimple与truncatedResult。完整静态模板尚未提供。"),
        ("4", "前置-用户问题拼接成上下文", "合并当前时间、用户问题、patientHistorySummary和完整附件文本。"),
        ("5", "意图分析", "用Qwen3.5-Plus输出一个或多个中医意图标签，每行一个。"),
        ("6", "过滤", "优先判断是否有文件；无文件时检查通识回答、非医学相关或其他意图。三个无文件出口实际汇入同一节点。"),
        ("7", "通过模型知识回答", "文件分支或证据不足回退时，使用Qwen3.8-Max及模型自身知识直接回答。"),
        ("8", "代码运行", "把多行意图文本转换成Array<string>，删除空行和严格等于“无”的行。"),
        ("9", "并行执行", "对每个意图运行内部检索流程；最大并发5，单轮报错重试3次。"),
        ("9.1", "开始", "在并行容器内提供当前数组下标与数组元素。"),
        ("9.2", "删除指令部分", "用正则删除两个##之间的文字；并非真正的指令识别。"),
        ("9.3", "检索词构建", "将当前意图和清理后的上下文改写成适合中医原始文献检索的短查询。"),
        ("9.4", "识别器", "把意图映射为允许检索的文献type代码。"),
        ("9.5", "检索", "在TCM知识库中执行all检索，最小相关性0.5，开启Rerank，并应用type筛选。"),
        ("9.6", "检索结果", "选择检索节点的搜索结果作为当前意图的批执行输出。"),
        ("10", "去重并重排序", "拍平成功结果、按dataId保留首条、清理snippet引用号、按rerankScore降序并重新编号。"),
        ("11", "知识引用-数组自定义引用模板", "将结构化证据转换为带[citation:index]、标题和来源信息的文本块。完整模板待确认。"),
        ("12", "通过知识库回答", "根据检索证据和病例生成专业中医回答，要求引用证据；同时要求可直接使用的辨证与加减处方。"),
        ("13", "记录知识库回答答案", "把知识库回答原文写入knowledgeSearchAnswer。"),
        ("14", "判断是否是有效回答", "只判断回答是否明确表示知识库无法回答，不校验医学正确性和引用真实性。"),
        ("15", "标记为二次回答", "无效回答分支将isReAnswer设为True，并转入模型自身知识回答。"),
        ("16", "记录回答节点答案", "正常知识库回答写入finalAnswer。"),
        ("17", "变量更新", "模型自身知识回答写入finalAnswer。"),
        ("18", "代码运行-生成输出数据集", "删除部分内部字段，生成准备保存的检索证据数组。"),
        ("19", "LYH-数据交换服务", "将证据保存为knowledgeSearch.json并返回访问URL。"),
        ("20", "整理输出最终回答", "使用Qwen3.5-Flash按用户格式要求整理finalAnswer，并强制删除[citation:n]；输出即最终答案。"),
    ]
    add_table(doc, ["序号", "准确节点名", "作用"], node_rows, [650, 2750, 5960], font_size=8.6)

    doc.add_heading("二、全局变量与数据传递", level=1)
    var_rows = [
        ("patientHistorySummary", "完整患者历史摘要；被放入<userQuery>内部。来源节点待确认。"),
        ("userAttachmentContexts", "附件/OCR对象数组；由附件拼接节点读取。来源与字段结构待确认。"),
        ("patientSummarySimple", "简化患者摘要；进入简易上下文。来源节点待确认。"),
        ("finalAnswer", "各回答分支汇合后的最终候选答案；最终格式节点读取。"),
        ("isReAnswer", "知识库回答被判无效时设为True。初始值与其他用途待确认。"),
        ("knowledgeSearchAnswer", "保存知识库回答原文；不等于最终发布内容。"),
        ("当前对话ID", "组成对象存储键。"),
        ("AI回复的ID", "组成对象存储键。"),
        ("当前时间", "注入完整上下文。"),
        ("用户问题 / 文件链接", "流程开始的外部输入。"),
    ]
    add_table(doc, ["变量", "传递方式与用途"], var_rows, [2700, 6660], font_size=9.2)

    doc.add_heading("三、按实际执行顺序拆解", level=1)
    doc.add_heading("1. 附件标准化", level=2)
    add_body(doc, "附件节点不执行OCR，而是处理上游已经生成的OCR对象。它支持数组或JSON字符串；整体异常时返回空字符串，单个附件异常时跳过。每个附件被包装为<attachment n>，完整结果用于完整上下文，去标签后的前300字符用于简易上下文。")
    add_bullet(doc, "通用能力：一次解析产生完整与简短两种输出，供不同下游复用。")
    add_bullet(doc, "写死内容：300字符上限、filename/context字段名、XML样式标签和OCR能力声明。")
    add_bullet(doc, "风险：除filename和context外的所有顶层属性都会写入上下文，可能夹带内部ID或不必要元数据。")

    doc.add_heading("2. 上下文组织", level=2)
    add_body(doc, "完整模板先写当前时间和系统/用户分界，再在<userQuery>中放入用户本轮问题与patientHistorySummary，最后追加完整附件文本。该结构帮助下游获得连续病史，但也把系统摘要与用户原话混合在同一标签内。")
    add_callout(doc, "来源混合", "模板声称<userQuery>只包含用户输入，实际还包含系统生成的患者历史摘要。下游无法单凭标签区分两者。", fill="FFF4E5", accent=GOLD)

    doc.add_heading("3. 意图识别与高层路由", level=2)
    add_body(doc, "意图分析定义通识回答、概念释义、鉴别诊断、方药查询、治法策略、医案经验、经典文献和针灸养生等标签，允许多选并要求一行一个。提示词还要求先对部分西医问题进行中医化。")
    add_bullet(doc, "非医学相关没有在正式意图列表中定义，却被过滤节点引用。")
    add_bullet(doc, "过滤使用“包含”而非“完全等于”，解释性文本也可能误触发。")
    add_bullet(doc, "通识回答、非医学相关和其他意图的出口均汇入同一代码节点，因此实际只区分“有文件”和“无文件”。")

    doc.add_heading("4. 多意图并行检索", level=2)
    add_body(doc, "无文件路径把意图文本转换成数组，并行容器对每个意图独立生成检索词、选择文献类型和执行知识库检索。该结构适合处理同时包含辨证、治法和方药等目标的复合问题。")
    add_bullet(doc, "最大并发数为5，单轮报错重试次数为3。")
    add_bullet(doc, "查询生成器默认输出1条，最多允许2条；但下游没有把两行再次拆成两个独立检索任务。")
    add_bullet(doc, "查询生成提示词禁止擅自辨证和添加方药，是降低错误传播的重要措施。")

    doc.add_heading("5. 意图到文献类型的映射", level=2)
    map_rows = [
        ("概念释义", "tcm_textbook, tcm_cw, tcm"),
        ("鉴别诊断", "tcm_tcmd, tcm_textbook, tcm_sd"),
        ("方药查询", "tcm_coffd, tcm_paf, tcm_hm, tcm_fr, tcm_textbook"),
        ("治法策略", "tcm_sd, tcm_coffd, tcm_textbook, tcm_fami"),
        ("医案经验", "tcm_sd, tcm_cw, tcm_mcd"),
        ("经典文献", "tcm_cw, tcm_fami"),
        ("针灸养生", "tcm_aam, tcm_fr"),
        ("未匹配标签", "tcm"),
    ]
    add_table(doc, ["意图", "允许检索的type代码"], map_rows, [2500, 6860], font_size=9.2)
    add_body(doc, "这些type代码的准确业务含义尚未提供。代码返回数组，但节点界面把extFilter声明为String；运行时是否被正确序列化必须通过节点调试确认。")

    doc.add_heading("6. 召回、去重与排序", level=2)
    add_body(doc, "检索节点使用TCM知识库、all检索、0.5最小相关性和Rerank；每个意图的成功结果在并行结束后被拍平。去重节点按dataId保留第一次出现的记录，然后清理snippet中的数字引用，按有效rerankScore降序排列，最后重新生成index。")
    add_bullet(doc, "同一dataId的后续高分记录可能在排序前被删除。")
    add_bullet(doc, "如果dataId代表整篇文献而非具体切片，可能丢失同文献中不同位置的有效片段。")
    add_bullet(doc, "不同查询产生的Rerank分数未必可以直接横向比较。")

    doc.add_heading("7. 循证回答与回退", level=2)
    add_body(doc, "知识引用模板把排序后的证据转换成带citation:index、标题和来源信息的文本；通过知识库回答节点再根据该文本与完整病例生成回答。提示词要求只依据检索结果，证据不足时停止，并为结论标注引用。")
    add_body(doc, "但回答节点没有直接接收结构化意图标签，只能重新猜测是否应使用“鉴别诊断/治法策略”的固定模板。其提示词还要求输出确定证型、完整病机、排除依据和可直接使用的加减处方，可能超过局部检索片段能够直接支持的范围。")
    add_callout(doc, "关键矛盾", "知识库回答正确地声明证据不足后，会被分类为“无效回答”，随后工作流转向模型自身知识继续生成医疗回答。安全停止因此被取消。", fill="FDECEC", accent=RED)

    doc.add_heading("8. 最终发布", level=2)
    add_body(doc, "正常知识库回答或模型知识回退答案都会写入finalAnswer。最终节点使用Qwen3.5-Flash根据用户格式要求整理文本，并强制删除所有[citation:n]。该节点之后没有医学正确性、引用一致性、隐私或内容差异检查。")

    doc.add_heading("四、防幻觉与错误传播措施", level=1)
    control_rows = [
        ("附件无文字警告", "附件拼接", "避免模型假装看懂未识别的图片；但不能判断OCR是否准确。"),
        ("固定意图标签", "意图分析", "减少自由输出；但下游未做标签白名单校验。"),
        ("禁止自行辨证", "检索词构建", "避免检索词把模型猜测固化为证据。"),
        ("按意图筛选文献type", "识别器/检索", "减少无关召回；映射错误也可能降低召回。"),
        ("向量+关键词+Rerank", "检索", "提高相关性；不能证明医学正确性。"),
        ("跨路去重与排序", "去重并重排序", "减少重复证据；去重策略可能丢失更高分或不同切片。"),
        ("证据不足允许停止", "通过知识库回答", "方向正确；但后续回退到模型自身知识。"),
        ("citation:index", "回答生成", "建立正文与证据编号关系；最终格式节点会删除。"),
        ("禁止格式节点新增医学信息", "整理输出最终回答", "降低排版改写风险；没有前后差异校验。"),
        ("对外导出前删除内部字段", "生成输出数据集", "减少内部信息暴露；只删除顶层字段且完整列表未知。"),
    ]
    add_table(doc, ["措施", "所在节点", "实际边界"], control_rows, [1900, 2150, 5310], font_size=8.8)

    doc.add_heading("五、通用能力与中医业务硬编码", level=1)
    add_table(
        doc,
        ["可迁移的通用能力", "写死的中医业务与技术配置"],
        [
            ("附件标准化；完整/简易上下文双输出；自然语言意图转数组；数组并发；查询改写；按意图路由数据源；检索/Rerank；结果聚合；引用模板；对象存储；全局状态汇合；最终格式渲染。",
             "8类中医意图；西医到中医映射；tcm_*文献类型；辨证/处方结构；病位规则；加减药30%规则；特定药物与方剂示例；300字符；并发5；重试3；相关性0.5；Token 20000；Rerank 10条。"),
        ],
        [4680, 4680],
        font_size=9.3,
        zebra=False,
    )

    doc.add_heading("六、值得肯定的设计", level=1)
    strengths = [
        "生成与检索分层：先生成检索式，再召回证据，再生成回答。",
        "复合问题拆成多个意图并发检索，可覆盖辨证、治法、方药等不同目标。",
        "按意图限制文献类型，减少无差别搜索产生的噪声。",
        "拍平、去重、排序和编号集中在统一聚合节点，职责清晰。",
        "证据数据与自然语言答案分开保存，有利于前端展示、调试和审计。",
        "知识库回答提示词允许证据不足，这是正确的安全方向。",
        "医学内容生成与最终格式整理分层，便于控制界面输出。",
    ]
    for item in strengths:
        add_bullet(doc, item)

    doc.add_heading("七、主要风险清单", level=1)
    risk_rows = [
        ("P0", "证据不足后回退模型知识", "取消安全停止，可能发布无检索证据的医疗回答。"),
        ("P0", "有效回答仅按措辞分类", "自信但错误、引用虚构或处方危险的回答仍会通过。"),
        ("P0", "要求可直接使用的个体化处方", "局部文献切片可能不足以支撑完整辨证与加减方案。"),
        ("P1", "最终删除citation角标", "正文结论失去与具体证据编号的可见对应。"),
        ("P1", "文件分支绕过检索", "任何带文件的问题直接使用模型自身知识。"),
        ("P1", "通识/非医学分支形同虚设", "两类问题仍进入中医检索。"),
        ("P1", "西医中文化规则冲突", "上游允许疾病替换，下游又明确禁止无条件替换。"),
        ("P1", "回答节点未收到结构化意图", "回答模板依赖模型重新猜测。"),
        ("P1", "extFilter类型不一致", "数组代码输出与String界面声明可能导致筛选失效。"),
        ("P1", "两条检索式未拆成两次检索", "多行查询可能被当成一个整体查询。"),
        ("P1", "按dataId先去重后排序", "可能保留较低分记录并删除更高分重复项。"),
        ("P1", "不同查询分数全局混排", "某一意图可能主导证据前排，削弱覆盖。"),
        ("P1", "患者摘要混入<userQuery>", "系统摘要与用户原话来源不清。"),
        ("P1", "提示注入隔离不足", "用户、OCR和文献片段都作为普通文本进入模型。"),
        ("P2", "最终格式模型无后置校验", "药名、限定语或表格对应关系被改写后直接发布。"),
        ("P2", "附件与导出清理不完整", "顶层之外的内部信息或敏感数据可能保留。"),
        ("P2", "静默错误处理", "附件或格式异常可能被跳过，运行结果缺少可诊断信息。"),
    ]
    risk_table = add_table(doc, ["级别", "问题", "潜在影响"], risk_rows, [720, 3300, 5340], font_size=8.8)
    for row in risk_table.rows[1:]:
        level = row.cells[0].text
        fill = "FDECEC" if level == "P0" else ("FFF4E5" if level == "P1" else LIGHT_GRAY)
        set_cell_shading(row.cells[0], fill)

    doc.add_heading("八、性能与维护风险", level=1)
    doc.add_heading("性能瓶颈", level=2)
    for text in [
        "每个意图都调用一次检索词生成模型，并执行一次检索和Rerank。",
        "最大5路并发且失败可重试3次，复杂问题会放大调用量。",
        "检索最大Token数为20000，多个意图合并后缺少最终证据数量上限。",
        "多个AI节点重复携带6轮聊天记录和文件链接。",
        "回答后还要调用分类模型和最终格式模型。",
    ]:
        add_bullet(doc, text)
    doc.add_heading("维护风险", level=2)
    for text in [
        "长提示词重复复制抗提示注入声明，修改时容易产生版本漂移。",
        "意图、文献type和处方规则分散在提示词与代码中，缺少单一配置源。",
        "内部节点ID与变量引用在复制、删除节点后容易失效。",
        "关键节点缺少严格JSON或枚举输出合同。",
        "节点名称可能高估能力，例如“有效回答”实际上只检测知识不足声明。",
    ]:
        add_bullet(doc, text)

    doc.add_heading("九、建议借鉴与避免照搬", level=1)
    do_text = "；".join([
        "拆分意图理解、检索词生成、召回和答案生成",
        "复合问题按意图并发检索",
        "按意图选择资料类型",
        "统一拍平、去重、排序和编号",
        "保留稳定证据定位信息",
        "把证据不足作为合法终态",
        "区分内部证据数据与外部展示数据",
        "把内容生成与格式渲染分层",
    ]) + "。"
    avoid_text = "；".join([
        "证据不足后回退模型自身知识",
        "最终删除全部引用却不提供替代追溯方式",
        "只靠大模型判断答案是否有效",
        "让模型生成可直接使用的个体化处方而缺少独立校验",
        "把系统摘要混入用户原话标签",
        "使用字符串包含作为关键分支条件",
        "不校验枚举标签和数据类型",
    ]) + "。"
    add_table(doc, ["值得借鉴", "不宜直接照搬"], [(do_text, avoid_text)], [4680, 4680], font_size=9.3, zebra=False)

    doc.add_heading("十、当前仍待确认的资料", level=1)
    pending = [
        "patientHistorySummary、patientSummarySimple和userAttachmentContexts的生成节点。",
        "前置-用户问题拼接成简易上下文的完整模板。",
        "知识引用-数组自定义引用模板的完整内容。",
        "代码运行-生成输出数据集的完整删除字段列表。",
        "所有tcm_*类型代码的准确业务含义。",
        "extFilter运行时究竟是字符串还是数组。",
        "并行执行失败结果、重试记录和错误出口。",
        "LYH访问URL的有效期、权限与前端使用方式。",
        "一次完整运行轨迹：意图、检索词、召回JSON、排序、引用回答和最终答案。",
        "文件模型是否真正支持图片、PDF或医学影像。",
        "最终界面是否在正文之外单独展示证据列表。",
    ]
    for item in pending:
        add_number(doc, item)

    conclusion_heading = doc.add_heading("总体结论", level=1)
    conclusion_heading.paragraph_format.space_before = Pt(6)
    conclusion_heading.paragraph_format.space_after = Pt(5)
    add_callout(
        doc,
        "结论",
        "检索工程思路完整，适合学习多意图RAG的分层结构；但模型知识回退、弱校验与最终去引用，使其不能被视为严格循证或可直接用于临床决策的系统。医疗场景必须把“证据不足”设为终止状态，而不是触发无证据回答。",
        fill="FDECEC",
        accent=RED,
        trailing_spacer=False,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
