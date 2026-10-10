from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_question_bank_docx as base


OUTPUT_PATH = Path(r"C:\Users\admin\Documents\ChatGPT\大查房\outputs\肿瘤大查房题目与答案解析.docx")


def add_cover(doc, cases):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(92)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("肿瘤大查房题目与答案解析")
    base.set_run_font(r, size=29, color=base.NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("单项选择题 · ABCD 选项 · 正确答案 · 答案解析")
    base.set_run_font(r, size=13, color=base.TEAL)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(4)
    rule.paragraph_format.space_after = Pt(22)
    base.set_paragraph_border(rule, color=base.TEAL, size="12")

    total_questions = sum(len(c["questions"].get("questions", [])) for c in cases)
    table = doc.add_table(rows=1, cols=3)
    values = [("病例", str(len(cases))), ("题目", str(total_questions)), ("日期", date.today().isoformat())]
    for idx, (label, value) in enumerate(values):
        cell = table.cell(0, idx)
        base.set_cell_shading(cell, base.LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label + "\n")
        base.set_run_font(r, size=8.5, color=base.MID_GRAY, bold=True)
        r = p.add_run(value)
        base.set_run_font(r, size=11, color=base.NAVY, bold=True)
    base.set_table_geometry(table, [3120, 3120, 3120], indent_dxa=0)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("仅用于医学教学与病例讨论，不替代临床诊疗决策")
    base.set_run_font(r, size=9, color=base.MID_GRAY, italic=True)
    doc.add_page_break()


def add_index(doc, cases):
    doc.add_heading("题目目录", level=1)
    table = doc.add_table(rows=1, cols=3)
    for i, text in enumerate(["病例编号", "病例标题", "题目数"]):
        cell = table.rows[0].cells[i]
        base.set_cell_shading(cell, base.NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        base.set_run_font(r, size=9.5, color=base.WHITE, bold=True)
    base.repeat_table_header(table.rows[0])
    for case in cases:
        detail = case["detail"]
        title = detail.get("overview", {}).get("hero", {}).get("title") or detail.get("home", {}).get("display_title", "")
        row = table.add_row()
        values = [case["case_no"], title, str(len(case["questions"].get("questions", [])))]
        for i, value in enumerate(values):
            p = row.cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            base.set_run_font(r, size=9.5, color=base.BLACK, bold=(i == 0))
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        base.prevent_row_split(row)
    base.set_table_geometry(table, [1400, 6760, 1200], indent_dxa=120)
    doc.add_page_break()


def add_question(doc, q, idx):
    heading = doc.add_heading(f"{q.get('id', f'Q{idx}')}｜{q.get('type', '')} · {q.get('phase', '')}", level=3)
    heading.paragraph_format.page_break_before = False

    timepoint = str(q.get("timepoint", "")).strip()
    if timepoint:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(timepoint)
        base.set_run_font(r, size=9, color=base.MID_GRAY, italic=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    r = p.add_run(str(q.get("question", "")))
    base.set_run_font(r, size=11, color=base.BLACK, bold=True)

    options = q.get("options", {})
    answer = str(q.get("answer", "")).upper()
    for letter in ("A", "B", "C", "D"):
        base.add_option(doc, letter, options.get(letter, ""), is_answer=(letter == answer))

    base.add_answer_box(doc, q)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("答案解析：")
    base.set_run_font(r, size=10.3, color=base.NAVY, bold=True)
    r = p.add_run(str(q.get("explain", "")))
    base.set_run_font(r, size=10.3, color=base.BLACK)

    divider = doc.add_paragraph()
    divider.paragraph_format.space_before = Pt(2)
    divider.paragraph_format.space_after = Pt(5)
    base.set_paragraph_border(divider, color=base.LINE_GRAY, size="4", space="2")


def build(cases):
    doc = Document()
    base.configure_styles(doc)
    for section in doc.sections:
        base.configure_page(section)
        base.setup_header_footer(section)
        section.header.paragraphs[0].clear()
        r = section.header.paragraphs[0].add_run("肿瘤大查房题目与答案解析")
        base.set_run_font(r, size=9, color=base.MID_GRAY, bold=True)

    add_cover(doc, cases)
    add_index(doc, cases)

    for case_idx, case in enumerate(cases, start=1):
        if case_idx > 1:
            doc.add_page_break()
        detail = case["detail"]
        title = detail.get("overview", {}).get("hero", {}).get("title") or detail.get("home", {}).get("display_title") or f"病例 {case['case_no']}"

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"病例 {case_idx:02d} · {case['case_no']}")
        base.set_run_font(r, size=9.5, color=base.TEAL, bold=True)
        h = doc.add_heading(title, level=1)
        h.paragraph_format.space_before = Pt(0)

        for idx, q in enumerate(case["questions"].get("questions", []), start=1):
            add_question(doc, q, idx)

    doc.core_properties.title = "肿瘤大查房题目与答案解析"
    doc.core_properties.subject = "单项选择题、正确答案及答案解析"
    doc.core_properties.author = "OncoRounds"
    doc.core_properties.keywords = "肿瘤, 大查房, 题目, 答案, 解析"
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


def main():
    cases = base.collect_cases()
    output = build(cases)
    print({
        "output": str(output),
        "case_count": len(cases),
        "question_count": sum(len(c["questions"].get("questions", [])) for c in cases),
    })


if __name__ == "__main__":
    main()
