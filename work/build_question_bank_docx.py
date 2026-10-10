import json
import os
import re
import hashlib
import urllib.request
from datetime import date
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE_DIRS = [
    Path(r"C:\Users\admin\Desktop\题库"),
    Path(r"C:\Users\admin\Desktop\题库2"),
]
WORK_DIR = Path(r"C:\Users\admin\Documents\ChatGPT\大查房\work\question_bank_docx")
IMAGE_DIR = WORK_DIR / "images"
OUTPUT_PATH = Path(r"C:\Users\admin\Documents\ChatGPT\大查房\outputs\肿瘤大查房题库汇编.docx")

CASE_RE = re.compile(r"^(\d+)-(病例详情页|题目JSON)\.json$")

NAVY = "17365D"
BLUE = "2E74B5"
TEAL = "167D8D"
LIGHT_BLUE = "E8F1F8"
LIGHT_TEAL = "E7F4F4"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "6B7280"
LINE_GRAY = "D9E1E8"
GREEN = "1E6B4E"
WHITE = "FFFFFF"
BLACK = "111827"


def load_json(path: Path):
    with path.open("r", encoding="utf-8-sig") as f:
        return json.load(f)


def collect_cases():
    grouped = {}
    for source_dir in SOURCE_DIRS:
        for path in sorted(source_dir.glob("*.json")):
            match = CASE_RE.match(path.name)
            if not match:
                continue
            case_no, kind = match.groups()
            grouped.setdefault(case_no, {})[kind] = path

    cases = []
    for case_no in sorted(grouped):
        pair = grouped[case_no]
        if "病例详情页" not in pair or "题目JSON" not in pair:
            continue
        detail = load_json(pair["病例详情页"])
        questions = load_json(pair["题目JSON"])
        cases.append({
            "case_no": case_no,
            "detail_path": pair["病例详情页"],
            "question_path": pair["题目JSON"],
            "detail": detail,
            "questions": questions,
        })
    if not cases:
        raise RuntimeError("题库目录中没有找到成对的病例详情页和题目 JSON。")
    return cases


def get_tab(detail, tab_id):
    for tab in detail.get("overview", {}).get("tabs", []):
        if tab.get("id") == tab_id:
            return tab
    return {}


def safe_filename_from_url(url):
    suffix = Path(url.split("?", 1)[0]).suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg", ".webp"}:
        suffix = ".img"
    digest = hashlib.sha256(url.encode("utf-8")).hexdigest()[:20]
    return digest + suffix


def download_images(cases):
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    url_to_path = {}
    failures = []
    for case in cases:
        images = get_tab(case["detail"], "images").get("image_items", [])
        for item in images:
            url = str(item.get("src", "")).strip()
            if not url or url in url_to_path:
                continue
            dest = IMAGE_DIR / safe_filename_from_url(url)
            if not dest.exists() or dest.stat().st_size == 0:
                req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
                try:
                    with urllib.request.urlopen(req, timeout=45) as response:
                        dest.write_bytes(response.read())
                except Exception as exc:
                    failures.append((url, str(exc)))
                    continue
            try:
                with Image.open(dest) as im:
                    im.verify()
                url_to_path[url] = dest
            except Exception as exc:
                failures.append((url, f"图片校验失败: {exc}"))
    return url_to_path, failures


def set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_border(paragraph, color=LINE_GRAY, size="8", space="4"):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=MID_GRAY)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 28, NAVY, 0, 10),
        ("Subtitle", 14, MID_GRAY, 0, 14),
        ("Heading 1", 18, NAVY, 18, 10),
        ("Heading 2", 14, BLUE, 14, 7),
        ("Heading 3", 12, TEAL, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Microsoft YaHei"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption.font.italic = False
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_together = True


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def setup_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("肿瘤大查房题库汇编")
    set_run_font(run, size=9, color=MID_GRAY, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(0)
    add_page_number(fp)


def add_cover(doc, cases):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(98)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("肿瘤大查房题库汇编")
    set_run_font(run, size=30, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run("病例详情 · 病程时间轴 · 影像资料 · 单项选择题与解析")
    set_run_font(run, size=13, color=TEAL)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(4)
    rule.paragraph_format.space_after = Pt(24)
    set_paragraph_border(rule, color=TEAL, size="12")

    total_questions = sum(len(c["questions"].get("questions", [])) for c in cases)
    meta = doc.add_table(rows=2, cols=3)
    values = [
        ("收录病例", str(len(cases))),
        ("题目总数", str(total_questions)),
        ("内容范围", "肺癌病例"),
        ("题量规则", "每例 2–5 题"),
        ("文档日期", date.today().isoformat()),
        ("资料类型", "脱敏教学资料"),
    ]
    for idx, (label, value) in enumerate(values):
        cell = meta.cell(idx // 3, idx % 3)
        set_cell_shading(cell, LIGHT_BLUE if idx < 3 else LIGHT_GRAY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label + "\n")
        set_run_font(r, size=8.5, color=MID_GRAY, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=11, color=NAVY, bold=True)
    set_table_geometry(meta, [3120, 3120, 3120], indent_dxa=0)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("仅用于医学教学与病例讨论，不替代临床诊疗决策")
    set_run_font(run, size=9, color=MID_GRAY, italic=True)
    doc.add_page_break()


def add_collection_index(doc, cases):
    doc.add_heading("收录病例", level=1)
    p = doc.add_paragraph("本汇编按病例编号排序。每章依次呈现病例概览、完整病程轴、病例影像及题目解析。")
    p.paragraph_format.space_after = Pt(10)

    table = doc.add_table(rows=1, cols=4)
    headers = ["编号", "病例标题", "题目", "图片"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=WHITE, bold=True)
    repeat_table_header(table.rows[0])

    for case in cases:
        detail = case["detail"]
        questions = case["questions"].get("questions", [])
        image_items = get_tab(detail, "images").get("image_items", [])
        title = detail.get("overview", {}).get("hero", {}).get("title") or detail.get("home", {}).get("display_title", "")
        row = table.add_row()
        values = [case["case_no"], title, str(len(questions)), str(len(image_items))]
        for i, value in enumerate(values):
            p = row.cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=9.5, color=BLACK, bold=(i == 0))
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        prevent_row_split(row)
    set_table_geometry(table, [1150, 6210, 1000, 1000], indent_dxa=120)
    doc.add_page_break()


def add_metadata_strip(doc, case):
    detail = case["detail"]
    questions = case["questions"]
    table = doc.add_table(rows=1, cols=4)
    values = [
        ("病例编号", detail.get("case_id") or questions.get("case_id") or case["case_no"]),
        ("病种", detail.get("cancer_type", "")),
        ("难度", questions.get("difficulty") or detail.get("difficulty", "")),
        ("题目数", str(len(questions.get("questions", [])))),
    ]
    for idx, (label, value) in enumerate(values):
        cell = table.cell(0, idx)
        set_cell_shading(cell, LIGHT_TEAL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label + "\n")
        set_run_font(r, size=8, color=MID_GRAY, bold=True)
        r = p.add_run(str(value))
        set_run_font(r, size=9.5, color=NAVY, bold=True)
    set_table_geometry(table, [2340, 2340, 2340, 2340], indent_dxa=0)


def add_label_value(doc, label, value):
    if value is None or str(value).strip() == "":
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    r = p.add_run(label + "：")
    set_run_font(r, size=10.5, color=NAVY, bold=True)
    r = p.add_run(str(value))
    set_run_font(r, size=10.5, color=BLACK)


def add_overview(doc, detail):
    doc.add_heading("病例概览", level=2)
    home = detail.get("home", {})
    if home.get("description"):
        p = doc.add_paragraph(home["description"])
        p.paragraph_format.space_after = Pt(9)
    summary_items = get_tab(detail, "summary").get("summary_items", [])
    for item in summary_items:
        add_label_value(doc, item.get("label", ""), item.get("value", ""))


def add_timeline(doc, detail):
    doc.add_heading("病程时间轴", level=2)
    items = get_tab(detail, "timeline").get("timeline_items", [])
    if not items:
        doc.add_paragraph("未提供结构化病程时间轴。")
        return
    table = doc.add_table(rows=1, cols=2)
    headers = ["时间", "诊疗事件"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=WHITE, bold=True)
    repeat_table_header(table.rows[0])
    for idx, item in enumerate(items):
        row = table.add_row()
        set_cell_shading(row.cells[0], LIGHT_BLUE if idx % 2 == 0 else "F7FAFC")
        set_cell_shading(row.cells[1], WHITE if idx % 2 == 0 else "FAFBFC")
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p0.add_run(str(item.get("date", "")))
        set_run_font(r, size=9, color=NAVY, bold=True)
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r = p1.add_run(str(item.get("title", "")))
        set_run_font(r, size=9.5, color=TEAL, bold=True)
        detail_text = str(item.get("detail", "")).strip()
        if detail_text:
            p2 = row.cells[1].add_paragraph()
            p2.paragraph_format.space_after = Pt(0)
            r = p2.add_run(detail_text)
            set_run_font(r, size=9.2, color=BLACK)
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        prevent_row_split(row)
    set_table_geometry(table, [1800, 7560], indent_dxa=120)


def add_image_alt(inline_shape, alt_text):
    try:
        doc_pr = inline_shape._inline.docPr
        doc_pr.set("descr", alt_text or "病例图片")
        doc_pr.set("title", alt_text or "病例图片")
    except Exception:
        pass


def image_size_for_doc(path, max_w=6.05, max_h=4.6):
    with Image.open(path) as im:
        w, h = im.size
    if w <= 0 or h <= 0:
        return Inches(max_w), None
    ratio = w / h
    width = max_w
    height = width / ratio
    if height > max_h:
        height = max_h
        width = height * ratio
    return Inches(width), Inches(height)


def add_images(doc, detail, url_to_path, download_failures):
    doc.add_heading("病例影像", level=2)
    items = get_tab(detail, "images").get("image_items", [])
    if not items:
        doc.add_paragraph("本病例未提供可用图片。")
        return
    for idx, item in enumerate(items, start=1):
        url = str(item.get("src", "")).strip()
        local_path = url_to_path.get(url)
        caption_text = item.get("caption") or item.get("alt") or f"病例图片 {idx}"
        if local_path and local_path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(5)
            width, height = image_size_for_doc(local_path)
            if height is None:
                shape = p.add_run().add_picture(str(local_path), width=width)
            else:
                shape = p.add_run().add_picture(str(local_path), width=width, height=height)
            add_image_alt(shape, str(item.get("alt") or caption_text))
            cap = doc.add_paragraph(style="Caption")
            cap.add_run(f"图 {idx}  {caption_text}")
            if item.get("timepoint"):
                cap.add_run(f"（{item['timepoint']}）")
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(f"图 {idx}  {caption_text}：图片下载失败，保留原始链接。")
            set_run_font(r, size=9.5, color="9B1C1C")
            p2 = doc.add_paragraph(url)
            p2.paragraph_format.space_after = Pt(8)
            for run in p2.runs:
                set_run_font(run, size=8.5, color=MID_GRAY)


def add_option(doc, letter, text, is_answer=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    r = p.add_run(f"{letter}. ")
    set_run_font(r, size=10.3, color=GREEN if is_answer else NAVY, bold=True)
    r = p.add_run(str(text))
    set_run_font(r, size=10.3, color=GREEN if is_answer else BLACK, bold=is_answer)


def add_answer_box(doc, question):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EAF6EF")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("正确答案：")
    set_run_font(r, size=10.5, color=GREEN, bold=True)
    r = p.add_run(f"{question.get('answer', '')}  {question.get('answer_text', '')}")
    set_run_font(r, size=10.5, color=GREEN, bold=True)
    set_table_geometry(table, [9360], indent_dxa=120)


def add_questions(doc, question_root):
    doc.add_heading("题目与答案解析", level=2)
    questions = question_root.get("questions", [])
    if not questions:
        doc.add_paragraph("本病例没有可发布题目。")
        return
    for idx, q in enumerate(questions, start=1):
        heading = doc.add_heading(f"{q.get('id', f'Q{idx}')}｜{q.get('type', '')} · {q.get('phase', '')}", level=3)
        heading.paragraph_format.page_break_before = False

        tp = doc.add_paragraph()
        tp.paragraph_format.space_after = Pt(4)
        r = tp.add_run(str(q.get("timepoint", "")))
        set_run_font(r, size=9, color=MID_GRAY, italic=True)

        stem = doc.add_paragraph()
        stem.paragraph_format.space_after = Pt(7)
        stem.paragraph_format.keep_together = True
        r = stem.add_run(str(q.get("question", "")))
        set_run_font(r, size=11, color=BLACK, bold=True)

        options = q.get("options", {})
        answer = str(q.get("answer", "")).upper()
        for letter in ("A", "B", "C", "D"):
            add_option(doc, letter, options.get(letter, ""), is_answer=(letter == answer))

        add_answer_box(doc, q)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run("答案解析：")
        set_run_font(r, size=10.3, color=NAVY, bold=True)
        r = p.add_run(str(q.get("explain", "")))
        set_run_font(r, size=10.3, color=BLACK)

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run("临床要点：")
        set_run_font(r, size=10, color=TEAL, bold=True)
        r = p.add_run(str(q.get("keyPoint", "")))
        set_run_font(r, size=10, color=BLACK)

        ref = q.get("guidelineRef", {}) or {}
        source_bits = [
            str(ref.get("source_title", "")).strip(),
            str(ref.get("source_version", "")).strip(),
            str(ref.get("source_section", "")).strip(),
            str(ref.get("authority", "")).strip(),
        ]
        source_line = "｜".join(bit for bit in source_bits if bit)
        if source_line or ref.get("evidence_summary"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run("证据来源：")
            set_run_font(r, size=8.8, color=MID_GRAY, bold=True)
            r = p.add_run(source_line)
            set_run_font(r, size=8.8, color=MID_GRAY)
            summary = str(ref.get("evidence_summary", "")).strip()
            if summary:
                p2 = doc.add_paragraph()
                p2.paragraph_format.left_indent = Inches(0.18)
                p2.paragraph_format.space_after = Pt(8)
                r = p2.add_run(summary)
                set_run_font(r, size=8.8, color=MID_GRAY, italic=True)

        divider = doc.add_paragraph()
        divider.paragraph_format.space_before = Pt(2)
        divider.paragraph_format.space_after = Pt(5)
        set_paragraph_border(divider, color=LINE_GRAY, size="4", space="2")


def build_doc(cases, url_to_path, download_failures):
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        setup_header_footer(section)

    add_cover(doc, cases)
    add_collection_index(doc, cases)

    for case_idx, case in enumerate(cases, start=1):
        if case_idx > 1:
            doc.add_page_break()
        detail = case["detail"]
        hero = detail.get("overview", {}).get("hero", {})
        title = hero.get("title") or detail.get("home", {}).get("display_title") or f"病例 {case['case_no']}"

        kicker = doc.add_paragraph()
        kicker.paragraph_format.space_before = Pt(0)
        kicker.paragraph_format.space_after = Pt(3)
        r = kicker.add_run(f"病例 {case_idx:02d} · {case['case_no']}")
        set_run_font(r, size=9.5, color=TEAL, bold=True)

        h = doc.add_heading(title, level=1)
        h.paragraph_format.space_before = Pt(0)
        if hero.get("subtitle"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(10)
            r = p.add_run(str(hero["subtitle"]))
            set_run_font(r, size=10.5, color=MID_GRAY, italic=True)

        add_metadata_strip(doc, case)
        add_overview(doc, detail)
        add_timeline(doc, detail)
        add_images(doc, detail, url_to_path, download_failures)
        add_questions(doc, case["questions"])

    doc.core_properties.title = "肿瘤大查房题库汇编"
    doc.core_properties.subject = "病例详情、时间轴、病例图片、单项选择题及答案解析"
    doc.core_properties.author = "OncoRounds"
    doc.core_properties.keywords = "肿瘤, 大查房, 病例, 题库, 时间轴, 影像"
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


def main():
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    cases = collect_cases()
    url_to_path, failures = download_images(cases)
    output = build_doc(cases, url_to_path, failures)
    report = {
        "output": str(output),
        "case_count": len(cases),
        "question_count": sum(len(c["questions"].get("questions", [])) for c in cases),
        "image_references": sum(len(get_tab(c["detail"], "images").get("image_items", [])) for c in cases),
        "downloaded_unique_images": len(url_to_path),
        "download_failures": failures,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
