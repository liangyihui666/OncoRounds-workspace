import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "1F4D78"
BLUE = "2E74B5"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7FA"
MID_GRAY = "667085"
LIGHT_BORDER = "D8E1EA"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_twips):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=11, bold=False, color=None, east_asia="Microsoft YaHei"):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def add_text(paragraph, text, size=11, bold=False, color=None):
    run = paragraph.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)
    return run


def style_paragraph(paragraph, before=0, after=6, line=1.25, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(paragraph, "第 ", size=8.5, color=MID_GRAY)
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_font(run, size=8.5, color=MID_GRAY)
    add_text(paragraph, " 页", size=8.5, color=MID_GRAY)


def normalize_options(options):
    if isinstance(options, dict):
        return [(str(k).upper(), str(v)) for k, v in options.items()]
    if isinstance(options, list):
        result = []
        for idx, value in enumerate(options):
            if isinstance(value, dict):
                key = value.get("key") or value.get("label") or chr(65 + idx)
                text = value.get("text") or value.get("value") or ""
            else:
                key = chr(65 + idx)
                text = value
            result.append((str(key).upper(), str(text)))
        return result
    return []


def source_name(question):
    ref = question.get("guidelineRef") or question.get("guideline_ref") or {}
    if isinstance(ref, dict):
        return ref.get("guidelineName") or ref.get("source_name") or ref.get("title") or ""
    return ""


def build(source_path, output_path):
    data = json.loads(Path(source_path).read_text(encoding="utf-8"))
    questions = data.get("questions", [])
    if len(questions) != 5:
        raise ValueError(f"Expected 5 questions, found {len(questions)}")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(header, "ONCOROUNDS  ·  模拟查房题集", size=8.5, bold=True, color=BLUE)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph()
    style_paragraph(title, after=5, keep=True)
    add_text(title, "ALK融合阳性肺腺鳞癌", size=22, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, after=8, keep=True)
    add_text(subtitle, "模拟查房 · 5道单项选择题", size=12, bold=True, color=BLUE)

    meta = doc.add_paragraph()
    style_paragraph(meta, after=10, keep=True)
    add_text(meta, "题目｜A–D选项｜参考答案｜简要解析", size=9.5, color=MID_GRAY)

    notice = doc.add_table(rows=1, cols=1)
    notice.alignment = WD_TABLE_ALIGNMENT.LEFT
    notice.autofit = False
    set_table_width(notice, 9360)
    set_cell_width(notice.cell(0, 0), 9360)
    set_cell_margins(notice.cell(0, 0), top=90, bottom=90, start=140, end=140)
    set_cell_shading(notice.cell(0, 0), PALE_GRAY)
    set_table_borders(notice, color="E3E8EF", size="4")
    p = notice.cell(0, 0).paragraphs[0]
    style_paragraph(p, after=0)
    add_text(p, "用途说明：", size=9, bold=True, color=NAVY)
    add_text(p, "仅用于医学教学与临床思维训练，不替代个体化诊疗决策。", size=9, color=MID_GRAY)

    for idx, question in enumerate(questions):
        if idx > 0:
            doc.add_page_break()
        else:
            spacer = doc.add_paragraph()
            style_paragraph(spacer, after=0)

        heading = doc.add_paragraph()
        style_paragraph(heading, before=5, after=4, keep=True)
        add_text(heading, f"{question.get('id', f'Q{idx+1}')}  ", size=16, bold=True, color=BLUE)
        add_text(heading, question.get("type", "单项选择题"), size=11, bold=True, color=NAVY)

        context = doc.add_paragraph()
        style_paragraph(context, after=9, keep=True)
        phase = question.get("phase", "")
        timepoint = question.get("timepoint", "")
        context_text = "  ·  ".join(x for x in (phase, timepoint) if x)
        add_text(context, context_text, size=9, color=MID_GRAY)

        stem = doc.add_paragraph()
        style_paragraph(stem, after=9, line=1.35, keep=True)
        add_text(stem, question.get("question", ""), size=11.5, bold=True, color="1D2939")

        options = normalize_options(question.get("options", []))
        opt_table = doc.add_table(rows=len(options), cols=2)
        opt_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        opt_table.autofit = False
        set_table_width(opt_table, 9360)
        set_table_borders(opt_table, color=LIGHT_BORDER, size="4")
        for row_idx, (label, text_value) in enumerate(options):
            label_cell, text_cell = opt_table.rows[row_idx].cells
            set_cell_width(label_cell, 620)
            set_cell_width(text_cell, 8740)
            set_cell_margins(label_cell, top=95, bottom=95, start=80, end=60)
            set_cell_margins(text_cell, top=95, bottom=95, start=120, end=120)
            label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(label_cell, PALE_BLUE)
            set_cell_shading(text_cell, WHITE if row_idx % 2 == 0 else "FAFBFC")
            lp = label_cell.paragraphs[0]
            lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(lp, after=0)
            add_text(lp, label, size=10.5, bold=True, color=BLUE)
            tp = text_cell.paragraphs[0]
            style_paragraph(tp, after=0, line=1.2)
            add_text(tp, text_value, size=10.5, color="1D2939")

        spacer2 = doc.add_paragraph()
        style_paragraph(spacer2, after=0)

        answer_table = doc.add_table(rows=1, cols=1)
        answer_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        answer_table.autofit = False
        set_table_width(answer_table, 9360)
        set_cell_width(answer_table.cell(0, 0), 9360)
        set_cell_margins(answer_table.cell(0, 0), top=140, bottom=140, start=160, end=160)
        set_cell_shading(answer_table.cell(0, 0), PALE_BLUE)
        set_table_borders(answer_table, color="AFCBE2", size="8")
        cell = answer_table.cell(0, 0)
        ap = cell.paragraphs[0]
        style_paragraph(ap, after=4)
        add_text(ap, "参考答案：", size=10.5, bold=True, color=NAVY)
        add_text(ap, question.get("answer", ""), size=12, bold=True, color=BLUE)
        answer_text = question.get("answer_text") or question.get("answerText") or ""
        if answer_text:
            add_text(ap, f"  {answer_text}", size=10.5, bold=True, color=NAVY)

        explanation = question.get("explain") or question.get("explanation") or ""
        if explanation:
            ep = cell.add_paragraph()
            style_paragraph(ep, after=4, line=1.25)
            add_text(ep, "解析：", size=9.5, bold=True, color=NAVY)
            add_text(ep, explanation, size=9.5, color="344054")

        key_point = question.get("keyPoint") or question.get("key_point") or ""
        if key_point:
            kp = cell.add_paragraph()
            style_paragraph(kp, after=0, line=1.2)
            add_text(kp, "教学要点：", size=9.5, bold=True, color=NAVY)
            add_text(kp, key_point, size=9.5, color="344054")

        src = source_name(question)
        if src:
            sp = doc.add_paragraph()
            style_paragraph(sp, before=5, after=0)
            add_text(sp, f"证据来源：{src}", size=8.5, color=MID_GRAY)

    props = doc.core_properties
    props.title = "ALK融合阳性肺腺鳞癌模拟查房5题"
    props.subject = "题目、选项、参考答案与解析"
    props.author = "OncoRounds"
    props.keywords = "ALK; 肺腺鳞癌; NSCLC; 模拟查房; 单项选择题"

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_alk_quiz_docx.py SOURCE_JSON OUTPUT_DOCX")
    build(sys.argv[1], sys.argv[2])
