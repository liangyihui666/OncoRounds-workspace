from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "竞品分析_临床决策App.md"
OUTPUT = ROOT / "竞品分析_临床决策App.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "202124"
MUTED = "6B7280"
LIGHT_BLUE = "F2F4F7"
CALLOUT = "F4F6F9"
RISK = "9B1C1C"
WHITE = "FFFFFF"


def set_run_font(run, size=11, bold=None, italic=None, color=INK):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row_props = row._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_hyperlink(paragraph, label, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.append(r_fonts)
    new_run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = label
    new_run.append(text)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\(https?://[^)]+\))")


def add_inline_markdown(paragraph, text, base_size=11, color=INK):
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=base_size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, color=DARK_BLUE)
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        else:
            m = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            add_hyperlink(paragraph, m.group(1), m.group(2))
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, color=color)


def add_para(doc, text, style=None, size=11, bold=False, italic=False, color=INK,
             align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, line=1.10,
             keep_with_next=False):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep_with_next
    if text:
        add_inline_markdown(p, text, base_size=size, color=color)
        if bold or italic:
            for run in p.runs:
                if bold:
                    run.bold = True
                if italic:
                    run.italic = True
    return p


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=9, color=MUTED)


def apply_base_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_picture(doc, image_path, alt_text):
    with Image.open(image_path) as img:
        width_px, height_px = img.size
        embed_path = image_path
        if image_path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}:
            asset_dir = ROOT / ".docx_assets"
            asset_dir.mkdir(exist_ok=True)
            embed_path = asset_dir / f"{image_path.stem}.png"
            if not embed_path.exists():
                img.convert("RGB").save(embed_path, "PNG", optimize=True)
    max_w, max_h = 5.85, 5.75
    scale = min(max_w / width_px, max_h / height_px)
    width_in = width_px * scale
    height_in = height_px * scale
    if width_in < 1.65:
        width_in = min(1.9, max_w)
        height_in = width_in * height_px / width_px
        if height_in > max_h:
            height_in = max_h
            width_in = height_in * width_px / height_px

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline_shape = run.add_picture(str(embed_path), width=Inches(width_in), height=Inches(height_in))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", alt_text)
    doc_pr.set("descr", alt_text)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_with_next = False
    run = cap.add_run(alt_text)
    set_run_font(run, size=9, italic=True, color=MUTED)


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.10
    add_inline_markdown(p, text, base_size=10, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        rows.append(cells)
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows, i


def widths_for_columns(cols):
    if cols == 1:
        return [9360]
    if cols == 2:
        return [2700, 6660]
    if cols == 3:
        return [2300, 3530, 3530]
    if cols == 4:
        return [1800, 2520, 2520, 2520]
    base = 9360 // cols
    widths = [base] * cols
    widths[-1] += 9360 - sum(widths)
    return widths


def add_markdown_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    set_table_geometry(table, widths_for_columns(cols))
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline_markdown(p, text, base_size=9.5, color=INK)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
                for run in p.runs:
                    run.bold = True
    set_repeat_table_header(table.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_title_block(doc):
    add_para(doc, "竞品分析", size=10, bold=True, color=BLUE, after=8)
    add_para(doc, "临床决策 App（丁香园）", size=28, bold=True, color=DARK_BLUE, after=6)
    add_para(doc, "产品体验、证据可信度与商业模式评估", size=14, color=MUTED, after=16)
    add_para(doc, "版本：2026 年 8 月 24 日｜基于用户实测截图与同一病例重复生成结果", size=9.5, color=MUTED, after=14)


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("临床决策 App 竞品分析")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    add_page_number(p)
    run = p.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_header_footer(section)
    apply_base_styles(doc)
    add_title_block(doc)

    i = 0
    skipped_source_title = False
    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("# ") and not skipped_source_title:
            skipped_source_title = True
            i += 1
            continue
        if stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            p_pr = p._p.get_or_add_pPr()
            p_bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "D7DBE2")
            p_bdr.append(bottom)
            p_pr.append(p_bdr)
            i += 1
            continue
        if stripped.startswith("## "):
            text = stripped[3:].strip()
            p = doc.add_paragraph(style="Heading 1")
            add_inline_markdown(p, text, base_size=16, color=BLUE)
            i += 1
            continue
        if stripped.startswith("### "):
            text = stripped[4:].strip()
            p = doc.add_paragraph(style="Heading 2")
            add_inline_markdown(p, text, base_size=13, color=BLUE)
            i += 1
            continue
        if stripped.startswith("#### "):
            text = stripped[5:].strip()
            p = doc.add_paragraph(style="Heading 3")
            add_inline_markdown(p, text, base_size=12, color=DARK_BLUE)
            i += 1
            continue
        if stripped.startswith("> "):
            callout_lines = []
            while i < len(lines) and lines[i].strip().startswith("> "):
                callout_lines.append(lines[i].strip()[2:])
                i += 1
            add_callout(doc, " ".join(callout_lines))
            continue
        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            alt, rel = image_match.groups()
            image_path = ROOT / rel
            if image_path.exists():
                add_picture(doc, image_path, alt or image_path.stem)
            else:
                add_callout(doc, f"图片缺失：{alt or rel}")
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_markdown_table(doc, rows)
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, bullet.group(1), base_size=11, color=INK)
            i += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline_markdown(p, numbered.group(1), base_size=11, color=INK)
            i += 1
            continue
        add_para(doc, stripped, size=11, after=6, line=1.10)
        i += 1

    props = doc.core_properties
    props.title = "临床决策 App（丁香园）竞品分析"
    props.subject = "产品体验、证据可信度与商业模式评估"
    props.author = ""
    props.last_modified_by = ""
    props.comments = "基于用户提供的产品截图与实测结果整理"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
