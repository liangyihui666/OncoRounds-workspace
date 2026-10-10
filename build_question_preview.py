from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = r"C:\Users\admin\Documents\ChatGPT\大查房\ALK阳性肺腺鳞癌_五题预览稿.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BORDER = "D7DEE8"


CASE_SUMMARY = (
    "57岁女性，初诊左肺下叶占位，临床分期cT2N1M0、IIB期。支气管镜有限活检示少量非小细胞癌伴坏死，"
    "形态及免疫组化倾向鳞状细胞癌；因组织量有限，未完成常规驱动基因检测。两周期术前化疗联合免疫治疗后"
    "原发灶增大，再次活检确诊肺腺鳞癌并检出ALK融合。患者接受阿来替尼后病灶及肺门淋巴结缩小、阻塞性肺不张复张，"
    "经MDT评估后接受左下肺叶切除及淋巴结清扫。术后继续阿来替尼。2024年12月颅脑MRI发现右侧颞叶新发约1.1cm"
    "环形强化灶，考虑转移可能性大；同期评估肺部控制良好、无系统进展，考虑中枢单灶进展，经MDT讨论后行射波刀并继续阿来替尼。"
)


QUESTIONS = [
    {
        "id": "Q1",
        "title": "有限活检标本的管理原则",
        "ability": "分子诊断",
        "phase": "初诊有限活检阶段",
        "timepoint": "2023-08-31，支气管镜活检标本进行病理评估时",
        "stem": "患者初诊支气管镜活检仅获得少量非小细胞癌组织并伴坏死，组织量有限，后续未能完成常规驱动基因检测。按照当前肺癌诊疗指南复盘，此类活检标本在病理诊断过程中最应遵循的标本管理原则是什么？",
        "options": [
            "先将全部标本用于形态学和免疫组化分型，分子检测可待复发后再考虑",
            "在完成病理诊断的同时，尽可能保留足够标本用于分子生物学和免疫治疗相关检测",
            "组织量有限时只需进行PD-L1检测，无需预留驱动基因检测标本",
            "只有确诊为晚期非鳞状NSCLC后，才需要考虑预留分子检测标本",
        ],
        "answer": "B",
        "answer_text": "在完成病理诊断的同时，尽可能保留足够标本用于分子生物学和免疫治疗相关检测",
        "evidence": "中华医学会肺癌临床诊疗指南(2025版)，doc_id 0e599051…，split 49",
    },
    {
        "id": "Q2",
        "title": "腺鳞癌的分子检测资格",
        "ability": "分子诊断",
        "phase": "再次活检与组织学重分类阶段",
        "timepoint": "2023年12月，再次活检确诊肺腺鳞癌后",
        "stem": "再次活检将患者的组织学诊断修正为肺腺鳞癌，后续手术标本也显示以鳞癌为主并含约10%腺癌成分。关于此类含腺癌成分NSCLC的分子检测资格，哪项最符合正式指南？",
        "options": [
            "只要鳞癌成分占优势，就不再需要进行驱动基因检测",
            "所有含腺癌成分的NSCLC，无论其他临床特征如何，均建议常规进行分子检测",
            "只有出现远处转移后，含腺癌成分的NSCLC才需要进行分子检测",
            "含腺癌成分的NSCLC只需检测EGFR，不需要检测ALK等其他驱动基因",
        ],
        "answer": "B",
        "answer_text": "所有含腺癌成分的NSCLC，无论其他临床特征如何，均建议常规进行分子检测",
        "evidence": "中国非小细胞肺癌放射治疗临床指南（2020版），doc_id cbed173d…，split 8",
    },
    {
        "id": "Q3",
        "title": "新辅助ALK-TKI研究证据的适用边界",
        "ability": "治疗决策",
        "phase": "ALK-TKI转化治疗与术前复评阶段",
        "timepoint": "2024-05-22，阿来替尼治疗后病灶缩小、准备评估手术机会时",
        "stem": "患者明确ALK融合阳性后接受阿来替尼，随后影像显示肿块和肺门淋巴结缩小、阻塞性肺不张复张，并进入手术评估。结合已经检索到的原始临床研究，关于ALK阳性可切除或潜在可切除NSCLC术前应用ALK-TKI的证据边界，哪项表述最严谨？",
        "options": [
            "该研究已经证明阿来替尼是免疫化疗失败后IIB期肺腺鳞癌的标准转化治疗",
            "研究中的客观缓解率较高，因此影像缩小本身即可决定立即手术，无需个体化评估",
            "一项12例回顾性研究显示新辅助ALK-TKI具有较高影像和病理缓解率，提示具有临床活性与可行性，但不足以单独确立标准方案或决定个体手术时机",
            "该研究是一项阿来替尼对比含铂化疗的随机III期试验，已证实新辅助ALK-TKI的生存优势",
        ],
        "answer": "C",
        "answer_text": "一项12例回顾性研究显示新辅助ALK-TKI具有较高影像和病理缓解率，提示具有临床活性与可行性，但不足以单独确立标准方案或决定个体手术时机",
        "evidence": "Targeted Oncology（2023），doc_id b8ababb7…，split 1；ORR 91.7%，MPR 75.0%，pCR 58.3%",
    },
    {
        "id": "Q4",
        "title": "术后ALK结果的临床定位",
        "ability": "治疗决策",
        "phase": "术后辅助治疗决策阶段",
        "timepoint": "2024-05-23，左下肺叶切除术后病理评估时",
        "stem": "患者术后病理记录为IB期，既往PCR已检出ALK融合，术后标本ALK(D5F3)亦为阳性。按照CACA肺癌指南复盘，关于这一ALK结果在术后管理中的定位，哪项最准确？",
        "options": [
            "ALK融合状态只在晚期或复发转移后有意义，术后无需纳入管理决策",
            "术后IB至IIIA期NSCLC的ALK融合检测结果用于指导辅助靶向治疗决策",
            "ALK阳性可以直接证明患者必须按固定疗程接受阿来替尼，无需考虑其他临床因素",
            "术后辅助治疗只需依据PD-L1结果，ALK融合状态不能提供决策信息",
        ],
        "answer": "B",
        "answer_text": "术后IB至IIIA期NSCLC的ALK融合检测结果用于指导辅助靶向治疗决策",
        "evidence": "CACA《肺癌》（2024），doc_id fd357ad4…，split 20",
    },
    {
        "id": "Q5",
        "title": "ALK-TKI期间中枢单灶进展的管理",
        "ability": "耐药与后线管理",
        "phase": "靶向治疗期间中枢单灶进展阶段",
        "timepoint": "当前教学复盘：回顾2024-12-16颅脑新发病灶的处理时点",
        "stem": "患者阿来替尼治疗期间颅脑MRI发现右侧颞叶新发约1.1cm环形强化灶，考虑转移可能性大；同期胸部肿瘤内科评估肺部控制良好、无系统进展，考虑中枢单灶进展。若按ALK-TKI治疗期间CNS进展进行当前教学复盘，最符合CSCO指南的管理策略是什么？",
        "options": [
            "继续原ALK-TKI并联合适当的局部治疗",
            "立即停用原ALK-TKI并常规更换为洛拉替尼，不再进行局部治疗",
            "立即停用所有系统治疗并进行全脑放疗",
            "继续原ALK-TKI单药观察，不对已出现的中枢病灶进行局部处理",
        ],
        "answer": "A",
        "answer_text": "继续原ALK-TKI并联合适当的局部治疗",
        "evidence": "CSCO非小细胞肺癌诊疗指南2026，doc_id 065ff4dc…，split 181",
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def create_option_numbering(document):
    numbering = document.part.numbering_part.element
    existing_abs = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(existing_abs, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "upperLetter")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)
    return abstract_id


def new_num_id(document, abstract_id):
    numbering = document.part.numbering_part.element
    existing = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_number(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# Resolved preset: compact_reference_guide. Header pattern: memo_masthead.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

meta_style = styles.add_style("Question Meta", WD_STYLE_TYPE.PARAGRAPH)
meta_style.font.name = "Calibri"
meta_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
meta_style.font.size = Pt(9.5)
meta_style.font.color.rgb = RGBColor.from_string(MUTED)
meta_style.paragraph_format.space_before = Pt(0)
meta_style.paragraph_format.space_after = Pt(8)
meta_style.paragraph_format.line_spacing = 1.15
meta_style.paragraph_format.keep_with_next = True

stem_style = styles.add_style("Question Stem", WD_STYLE_TYPE.PARAGRAPH)
stem_style.font.name = "Calibri"
stem_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
stem_style.font.size = Pt(11.5)
stem_style.font.bold = True
stem_style.font.color.rgb = RGBColor.from_string(INK)
stem_style.paragraph_format.space_before = Pt(4)
stem_style.paragraph_format.space_after = Pt(10)
stem_style.paragraph_format.line_spacing = 1.3
stem_style.paragraph_format.keep_with_next = True

header_p = section.header.paragraphs[0]
header_p.text = "肿瘤大查房题目预览｜ALK阳性肺腺鳞癌"
header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
header_p.paragraph_format.space_after = Pt(0)
for run in header_p.runs:
    set_run_font(run, size=9, color=MUTED)
footer_p = section.footer.paragraphs[0]
footer_p.add_run("内部审阅稿  ·  ")
for run in footer_p.runs:
    set_run_font(run, size=9, color=MUTED)
add_page_number(footer_p)

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(8)
title.paragraph_format.space_after = Pt(4)
run = title.add_run("ALK阳性肺腺鳞癌全程管理")
set_run_font(run, size=24, bold=True, color=DARK_BLUE)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
run = subtitle.add_run("五道单选题预览稿  ·  进阶难度  ·  内部审阅")
set_run_font(run, size=11, color=MUTED)

intro_table = doc.add_table(rows=1, cols=1)
set_table_geometry(intro_table, [9360])
mark_header_row(intro_table.rows[0])
cell = intro_table.cell(0, 0)
set_cell_shading(cell, LIGHT_BLUE)
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(4)
r = p.add_run("病例摘要")
set_run_font(r, size=11, bold=True, color=DARK_BLUE)
p = cell.add_paragraph(CASE_SUMMARY)
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.25

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(10)
note.paragraph_format.space_after = Pt(0)
r = note.add_run("审阅说明：")
set_run_font(r, size=10, bold=True, color=DARK_BLUE)
r = note.add_run("正文仅展示题干与选项；答案和证据定位集中放在文末，便于先独立判断题目质量。")
set_run_font(r, size=10, color=MUTED)

option_abstract = create_option_numbering(doc)

for q in QUESTIONS:
    heading = doc.add_paragraph(style="Heading 1")
    heading.paragraph_format.page_break_before = True
    heading.add_run(f"{q['id']}｜{q['title']}")

    meta = doc.add_paragraph(style="Question Meta")
    meta.add_run(f"能力维度：{q['ability']}    阶段：{q['phase']}\n时间点：{q['timepoint']}")

    stem = doc.add_paragraph(q["stem"], style="Question Stem")

    num_id = new_num_id(doc, option_abstract)
    for option in q["options"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        apply_number(p, num_id)
        p.add_run(option)

    review = doc.add_paragraph()
    review.paragraph_format.space_before = Pt(14)
    review.paragraph_format.space_after = Pt(0)
    r = review.add_run("审阅记录：")
    set_run_font(r, size=9.5, bold=True, color=MUTED)
    r = review.add_run("  □题干清楚  □选项唯一  □难度合适  □证据匹配")
    set_run_font(r, size=9.5, color=MUTED)

answer_heading = doc.add_paragraph(style="Heading 1")
answer_heading.paragraph_format.page_break_before = True
answer_heading.add_run("答案与证据定位")

answer_note = doc.add_paragraph("本页供审阅核对使用，不建议在正式答题界面中与题目同时展示。")
answer_note.paragraph_format.space_after = Pt(10)
for run in answer_note.runs:
    set_run_font(run, size=10, color=MUTED, italic=True)

table = doc.add_table(rows=1, cols=4)
set_table_geometry(table, [720, 720, 4260, 3660])
mark_header_row(table.rows[0])
headers = ["题号", "答案", "正确选项", "主要证据"]
for idx, text in enumerate(headers):
    cell = table.rows[0].cells[idx]
    set_cell_shading(cell, LIGHT_BLUE)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run_font(r, size=9.5, bold=True, color=DARK_BLUE)

for q in QUESTIONS:
    row = table.add_row()
    values = [q["id"], q["answer"], q["answer_text"], q["evidence"]]
    for idx, value in enumerate(values):
        cell = row.cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(value)
        set_run_font(r, size=9.2, bold=(idx == 1), color=INK)

set_table_geometry(table, [720, 720, 4260, 3660])

doc.core_properties.title = "ALK阳性肺腺鳞癌全程管理五道题预览稿"
doc.core_properties.subject = "肿瘤大查房题目与选项审阅"
doc.core_properties.author = ""
doc.core_properties.keywords = "ALK, 肺腺鳞癌, 肿瘤大查房, 单选题"
doc.save(OUTPUT)
print(OUTPUT)
