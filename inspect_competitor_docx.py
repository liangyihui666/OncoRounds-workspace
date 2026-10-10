from pathlib import Path
import sys
from docx import Document

SRC = Path(r"C:\Users\admin\Downloads\临床决策 App 竞品分析 (1).docx")
doc = Document(SRC)
start = int(sys.argv[1]) if len(sys.argv) > 1 else 0
end = int(sys.argv[2]) if len(sys.argv) > 2 else len(doc.paragraphs)
show_tables = "--tables" in sys.argv
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)} range={start}:{end}")
for i, p in enumerate(doc.paragraphs[start:end], start=start):
    text = p.text.replace("\n", "\\n")
    if text.strip() or p._p.xpath('.//w:drawing'):
        image = " [IMAGE]" if p._p.xpath('.//w:drawing') else ""
        style_name = p.style.name if p.style is not None else "(none)"
        print(f"P{i:03d}\t{style_name!r}\t{text}{image}")
if show_tables:
    for ti, table in enumerate(doc.tables):
        print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
        for ri, row in enumerate(table.rows):
            print("  R%02d\t%s" % (ri, " | ".join(c.text.replace("\n", " / ") for c in row.cells)))
