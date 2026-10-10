from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SRC = Path(r"C:\Users\admin\Downloads\临床决策 App 竞品分析 (1).docx")
OUT_DIR = Path(r"C:\Users\admin\Documents\ChatGPT\大查房\竞品分析_临床决策")
OUT = OUT_DIR / "临床决策_App_竞品分析_精简版.docx"


def paragraph_by_exact(doc, text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise KeyError(f"Paragraph not found: {text}")


def paragraph_by_prefix(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(f"Paragraph prefix not found: {prefix}")


def set_text(p, text):
    first_rpr = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for child in list(p._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            p._p.remove(child)
    run = p.add_run(text)
    if first_rpr is not None:
        run._r.insert(0, first_rpr)
    return p


def remove_paragraph(p):
    parent = p._element.getparent()
    parent.remove(p._element)
    p._p = p._element = None


def insert_after_paragraph(p, text, template=None):
    new_p = OxmlElement("w:p")
    if template is not None and template._p.pPr is not None:
        new_p.append(deepcopy(template._p.pPr))
    p._p.addnext(new_p)
    para = Paragraph(new_p, p._parent)
    run = para.add_run(text)
    if template is not None and template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    return para


def insert_after_table(table, text, template):
    new_p = OxmlElement("w:p")
    if template._p.pPr is not None:
        new_p.append(deepcopy(template._p.pPr))
    table._tbl.addnext(new_p)
    para = Paragraph(new_p, table._parent)
    run = para.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    return para


def add_index_table(doc, after_p):
    title = insert_after_paragraph(after_p, "问题索引（完整论证与证据定位）", paragraph_by_exact(doc, "P0-1：隐私与安全承诺过度"))
    table = doc.add_table(rows=1, cols=3)
    table.style = doc.tables[9].style
    table.autofit = True
    headers = ["编号", "问题", "证据"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
    rows = [
        ("P0-1", "隐私与安全承诺过度", "Step 5-B"),
        ("P0-2", "证据可信度展示偏营销", "Step 8-A / 10"),
        ("P1-1", "病例核对效率不足", "Step 6"),
        ("P1-2", "图片失败恢复不行动", "Step 5-A"),
        ("P1-3", "删除范围不透明", "Step 4 / 4-A / 4-D"),
        ("P1-4", "专业认证治理说明不足", "Step 4-C"),
        ("P1-5", "分析过程暴露内部思考", "Step 7"),
        ("P2-1", "移动端长内容可读性", "Step 6 / 六章"),
        ("P2-2", "历史资产管理偏基础", "Step 4"),
    ]
    for number, issue, evidence in rows:
        cells = table.add_row().cells
        cells[0].text = number
        cells[1].text = issue
        cells[2].text = evidence
    title._p.addnext(table._tbl)
    return table


def add_evidence_line(doc, heading_text, evidence_text):
    heading = paragraph_by_exact(doc, heading_text)
    body = heading._p.getnext()
    body_para = Paragraph(body, heading._parent) if body is not None and body.tag.endswith("}p") else None
    return insert_after_paragraph(heading, f"证据：{evidence_text}", body_para or heading)


def table_by_first_row(doc, expected):
    for table in doc.tables:
        values = [c.text.strip() for c in table.rows[0].cells]
        if values == expected:
            return table
    raise KeyError(f"Table not found: {expected}")


def table_containing(doc, needle):
    for table in doc.tables:
        if needle in table.cell(0, 0).text:
            return table
    raise KeyError(f"Table containing text not found: {needle}")


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)

    # 三、体验流程：仅保留可见事实、健康度与短交叉引用。
    replacements = {
        "历史会话支持单条重命名和删除，这是基本的数据整理能力；未见搜索、筛选、日期、病例标签、收藏、归档或批量操作，当会话量增大时可管理性可能迅速下降。":
            "历史会话支持单条重命名和删除；未见搜索、筛选、日期、病例标签、收藏、归档或批量操作。→ 详见 P2-2。删除后是否同步清除服务器副本及原图，证据不足。→ 详见 P1-3。",
        "仍未确认“个人信息下载”是否包含病例、会话、上传原图和生成内容；单条会话删除是否同步删除服务器副本、缓存和原图也未知。":
            "个人信息下载是否包含病例、会话、上传原图和生成内容尚未确认。→ 详见 P1-3。",
        "页面宣称认证信息“仅用于资格认证”，但截图未说明审核主体、核验数据源、是否需要执业证件、材料保存期限、认证撤回、失败后的数据处理及访问权限。":
            "页面宣称认证信息“仅用于资格认证”，但未展示审核主体、核验数据源、材料保存期限、认证撤回及访问权限。→ 详见 P1-4。",
        "“95%的医疗同行已通过认证”属于强社会认同文案，其统计口径、样本范围和更新时间不可见，若无法解释容易成为不可验证的转化话术。":
            "页面显示“95%的医疗同行已通过认证”，未见统计口径、样本范围和更新时间。→ 详见 P1-4。",
        "提交按钮在没有输入时保持禁用，状态反馈清楚；但页面继续使用“不会识别任何隐私信息”的绝对承诺，和真实病历输入风险不匹配。":
            "提交按钮在没有输入时保持禁用；页面仍使用“不会识别任何隐私信息”文案。→ 详见 P0-1。",
        "历史会话支持重命名和删除，用户可以整理或移除单条任务；用户实测删除后不可恢复。是否存在二次确认，以及删除是否连带清除服务器原图、结构化数据、模型记录和缓存，仍无法确认。":
            "历史会话支持重命名和删除；用户实测删除后不可恢复。删除是否连带清除原图、结构化数据、模型记录和缓存尚未确认。→ 详见 P1-3。",
        "“不会识别患者和医生任何隐私信息”是高风险承诺。仅凭截图无法证明自动脱敏能力；如果图片含姓名、住院号、二维码或日期，用户可能产生错误安全感。":
            "页面仍显示“不会识别患者和医生任何隐私信息”。→ 详见 P0-1。",
        "提交前未见可视化脱敏检查、敏感字段预览、确认勾选或数据保存期限说明。":
            "提交前未见可视化脱敏检查。→ 详见 P0-1。",
        "失败原因仍然笼统：没有说明是图片模糊、文字过小、裁切不完整、反光、格式不支持，还是网络/OCR服务异常，因此用户只能盲目重试。":
            "页面未说明解析失败的具体原因。→ 详见 P1-2。",
        "上传前没有可见的图片质量预检、裁切/旋转、清晰度提示或示例对照，导致错误在提交后才暴露。":
            "上传前未见图片质量预检、裁切/旋转或清晰度提示。→ 详见 P1-2。",
        "因此，“不会识别患者和医生任何隐私信息”的绝对表述与可见产品机制不匹配。更准确的承诺应区分：是否识别、是否传输、是否用于模型处理、是否在输出展示、保存多久、由谁可访问、如何删除。":
            "→ 详见 P0-1。",
        "两张原始病历对人眼具有较高可读性，却曾出现笼统“解析失败”提示。无法据此确定真实原因；可能与版面倾斜/弯曲、正文占比、页面分割、低密度页或 OCR 服务异常有关。产品不应把所有失败都归因于“请上传清晰图片”。":
            "两张原始病历曾触发笼统“解析失败”提示。→ 详见 P1-2。",
        "长文本以单列连续排布，缺少时间轴、异常值突出、信息缺失提示和关键矛盾提示；医生核对成本仍高。":
            "长文本以单列连续排布，缺少时间轴、异常值突出、信息缺失提示和关键矛盾提示。→ 详见 P1-1。",
        "当前界面暴露了大量模型生成的中间问题，文字冗长且重复“分析完成”；这更像调试日志，不像临床工作摘要。":
            "当前界面暴露大量中间问题，并重复显示“分析完成”，呈现形态接近调试日志。→ 详见 P1-5。",
        "医疗 AI 更需要明确显示“哪些证据支持哪条结论、哪些问题证据不足、何时停止给出建议”，而不只是评议次数。":
            "→ 详见 P0-2。",
        "引用元数据仍有瑕疵：部分标题被截断，病例文本还出现由字符串自动识别造成的异常链接；个别句子同时引用多个来源，但读者仍需自行判断每个来源究竟支持哪一部分。":
            "输出中出现引用错位和 [ccitation:4] 占位符泄露；部分标题被截断，个别句子同时引用多个来源。→ 详见 P0-2。",
        "“100% 权威证据库”“100% 真实可溯源”是付费价值主张，但已提供的病例输出出现引用错位和 [ccitation:4] 占位符泄露。这形成明显的“付费承诺—实际质量”落差，应作为竞品风险重点记录。":
            "会员页使用“100% 权威证据库”“100% 真实可溯源”；该口号与 Step 8-A 的引用错位和占位符泄露形成落差。→ 详见 P0-2。",
    }
    for old, new in replacements.items():
        set_text(paragraph_by_exact(doc, old), new)

    for text in [
        "截图无法确认失败是否消耗每日额度、原图是否继续保存、用户已填写的补充信息是否保留，以及重试次数是否有限制。",
        "上传界面未见敏感字段自动框选、遮挡预览、提交前阻断、处理地域/第三方说明、保存期限或删除入口。对真实医疗资料而言，这是 P0 级信任缺口。",
        "“修改病例”按钮对关键安全动作的视觉权重低于“开始分析”，可能诱导用户跳过核对。",
        "因此，该输出适合作为医生复核的分析草稿，不应被产品包装成已经完成独立验证的最终临床方案。本报告不对其中具体诊疗结论作医学背书。",
    ]:
        remove_paragraph(paragraph_by_exact(doc, text))

    # Step 8-A 只有用户提供的一个版本，不使用“最新版/前一版”表述。
    set_text(paragraph_by_prefix(doc, "Step 8-A："), "Step 8-A：病例分析输出质量 — 健康度：一般")
    set_text(paragraph_by_prefix(doc, "Step 4-C："), "Step 4-C：专业身份认证 — 健康度：一般")
    set_text(paragraph_by_prefix(doc, "Step 10："), "Step 10：选择会员套餐并支付 — 健康度：良好")

    # 四-A：商业启示移入七章；此处只保留商业模式事实。
    for text in [
        "对 OncoRounds 的启示",
        "不必复制“通用医学工具大礼包”；更适合按病例或题组产出价值收费，例如科室版病例额度、验证题组包、证据更新订阅和团队审核席位。",
        "免费或低价层可以提供病例结构化预览，但“完整五题 + 独立验证 + 题目级证据定位 + 导出”可作为核心付费权益。",
        "若采用自动续费，应让首购价、续费价、额度和取消路径同屏可见，避免把商业转化建立在价格理解成本上。",
        "不要使用“100% 权威”“100% 可溯源”等绝对化口号；应展示可验证的通过率、引用完整率和发布门禁。",
    ]:
        remove_paragraph(paragraph_by_exact(doc, text))

    # 五章：统一编号并增加问题索引与证据定位。
    heading_map = {
        "P0：隐私与安全承诺过度": ("P0-1：隐私与安全承诺过度", "Step 5-B"),
        "P0：证据可信度的展示仍偏营销": ("P0-2：证据可信度展示偏营销", "Step 8-A / Step 10"),
        "P1：病例核对效率不足": ("P1-1：病例核对效率不足", "Step 6"),
        "P1：病例图片失败恢复不够可行动": ("P1-2：图片失败恢复不行动", "Step 5-A"),
        "P1：账号与病例数据控制已有入口，但删除范围不透明": ("P1-3：删除范围不透明", "Step 4 / Step 4-A / Step 4-D"),
        "P1：专业认证采集高敏感信息，治理说明不足": ("P1-4：专业认证治理说明不足", "Step 4-C"),
        "P1：分析过程暴露过多内部思考": ("P1-5：分析过程暴露内部思考", "Step 7"),
        "P2：移动端长内容可读性一般": ("P2-1：移动端长内容可读性", "Step 6 / 六章"),
        "P2：历史资产管理仍偏基础": ("P2-2：历史资产管理偏基础", "Step 4"),
    }
    new_headings = []
    for old, (new, evidence) in heading_map.items():
        set_text(paragraph_by_exact(doc, old), new)
        new_headings.append((new, evidence))
    add_index_table(doc, paragraph_by_exact(doc, "五、主要短板与风险"))
    for heading, evidence in new_headings:
        add_evidence_line(doc, heading, evidence)

    # 七章：只留不要照搬、商业策略和对九章的指向。
    for text in [
        "应借鉴并升级的部分",
        "借鉴病例图片上传，但增加脱敏门禁。 上传前提供脱敏清单；OCR 后将姓名、住院号、电话、二维码、精确日期等标为待确认字段。",
        "更进一步，应在本地或可信边界内先生成不可逆遮挡副本，默认只上传副本；同时增加图片质量预检，提交前检查清晰度、倾斜、弯曲、反光、正文占比、页码完整性和文档分割。失败时保留已填内容、给出具体原因，并明确不消耗病例额度。",
        "借鉴结构化病例确认，但改为肿瘤专用结构。 采用诊断、分期、病理、分子、既往线次、疗效、不良反应、当前决策点的时间轴，并保留原图对照。",
        "借鉴分析阶段可视化，但显示验证状态。 用“病例要素完整性 → 指南检索 → 临床研究检索 → 生成五题 → 独立验证 → 发布门禁”替代泛化的思考日志。",
        "借鉴引用原文弹层，并做到题目级证据映射。 每道题显示来源类型、标题、年份、定位、支持的答案点和局限。",
        "把证据不足做成正式结果。 清楚显示 PASS / REVISE / STOP_INSUFFICIENT，证据不足时宁可少发布或停止，也不凑满五题。",
        "用教学留存替代资讯流留存。 可做“本周病例复盘、错题回看、专科题组订阅、证据更新提醒”，比泛医学资讯更符合大查房定位。",
        "把专业身份与数据控制做成基础设施。 可借鉴职业分流认证与隐私中心，但认证信息应最小化采集，并显示字段用途与保存期限；同时为病例、题组、原图和证据缓存提供查看、导出、删除、保留期限与审计记录。",
    ]:
        remove_paragraph(paragraph_by_exact(doc, text))
    anchor = paragraph_by_exact(doc, "不要对隐私做绝对承诺，也不要在未完成证据验证时发布看似确定的结论。")
    heading_template = paragraph_by_exact(doc, "不要照搬的部分")
    body_template = paragraph_by_exact(doc, "不要做“病例、处方、指南什么都能问”的通用首页，容易稀释肿瘤教学定位。")
    anchor = insert_after_paragraph(anchor, "商业策略", heading_template)
    for text in [
        "按病例或题组产出价值收费，例如科室病例额度、验证题组包、证据更新订阅和团队审核席位。",
        "免费或低价层提供病例结构化预览；完整五题、独立验证、题目级证据定位与导出作为核心付费权益。",
        "首购价、续费价、额度和取消路径同屏可见。",
        "具体行动与优先级见九章表格。",
    ]:
        anchor = insert_after_paragraph(anchor, text, body_template)

    # 八章：只保留定位句（原有表格）与指向九章的交叉引用。
    for text in [
        "支撑这一主张的四项产品证据应在界面上直接可见：",
        "固定五题及难度/知识点分布；",
        "病例事实、指南、原创临床研究的证据配额；",
        "每题证据定位与答案理由；",
        "独立验证结果和发布门禁。",
    ]:
        remove_paragraph(paragraph_by_exact(doc, text))
    positioning_table = table_containing(doc, "临床决策帮助医生")
    insert_after_table(positioning_table, "证据可见性要求见九章 P0。", paragraph_by_exact(doc, "八、建议的差异化产品主张"))

    # 九章：唯一行动表，补充续费透明度。
    action_table = table_by_first_row(doc, ["优先级", "建议", "目标"])
    row = action_table.add_row().cells
    row[0].text = "P1"
    row[1].text = "首购价、续费价、额度与取消路径同屏可见"
    row[2].text = "降低续费理解成本与自动续费争议"

    # 统一核心交叉引用句式。
    set_text(paragraph_by_exact(doc, "核心事实发生漂移：原始病例列出了 5、6、7、10、11、12、14 区淋巴结，其中 5 区 1/1 阳性，其余送检区域阴性；本次生成却写成“仅报告 5 区淋巴结 1 枚送检”，并据此判断未达到采样标准。这不是措辞差异，而是对输入事实的漏读，直接影响分期充分性、PORT 和后续治疗判断。"),
             "同一病例重复生成时，原始病例所列 5、6、7、10、11、12、14 区淋巴结被改写为“仅报告 5 区 1 枚送检”，核心事实发生漂移。")
    set_text(paragraph_by_exact(doc, "因此，重复生成测试说明该产品的主要风险已从“能否找到文献”转向“病例事实是否稳定、结论是否经过发布前一致性校验”。即使每条引用都可正确跳转，也不能弥补输入事实漏读和版本间冲突。"),
             "重复生成出现事实漂移、结论冲突和发布前质检残留。→ 详见 P0-2。")

    doc.core_properties.title = "临床决策 App 竞品分析（精简版）"
    doc.core_properties.subject = "按 Step 事实—五章结论—九章行动的单一归属规则精简"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
